from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE-Chuong-1-2.docx"
DIAGRAM = ROOT / "docs" / "architecture" / "jpulse-container-main.png"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-chapter-3-partial.docx"

doc = Document(SOURCE)
headings_before = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]


def heading(text: str):
    matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == text]
    assert len(matches) == 1, (text, len(matches))
    return matches[0]


def add_before(anchor, text: str, lead: str | None = None):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    if lead:
        run = paragraph.add_run(lead)
        run.bold = True
    paragraph.add_run(text)


# 3.1. Nguyên tắc thiết kế
anchor = heading("Kiểu kiến trúc được lựa chọn")
add_before(anchor, "JPULSE ưu tiên phân tách trách nhiệm và bảo toàn tính đúng đắn của dữ liệu nghiệp vụ. Những nguyên tắc sau rút ra từ ranh giới hệ thống, ràng buộc kỹ thuật và yêu cầu chất lượng đã nêu trong SRS (Mục 2.4–2.6; Chương 7).")
add_before(anchor, "Giao diện web trình bày trạng thái và tiếp nhận thao tác; API thực thi quy tắc nghiệp vụ, phân quyền, chuyển trạng thái và điều phối tích hợp; Firestore và Storage lưu trạng thái bền vững. Các thay đổi nghiệp vụ quan trọng phải đi qua API để có cùng một điểm kiểm soát quyền, kiểm tra đầu vào và ghi audit.", "Phân tách trách nhiệm: ")
add_before(anchor, "Mọi truy cập dữ liệu và lệnh xử lý phải gắn với định danh, quyền theo vai trò và phạm vi cơ sở. Khi thiếu hoặc sai ngữ cảnh quyền, hệ thống từ chối mặc định. Đường truy cập trực tiếp từ web đến Firestore/Storage chỉ được dùng trong phạm vi được Firebase Security Rules và chính sách lưu trữ cho phép.", "Quyền tối thiểu theo phạm vi: ")
add_before(anchor, "JPULSE sở hữu trạng thái quy trình và dữ liệu nội bộ; hệ thống ngoài giữ quyền đối với kết quả do họ phát hành. Lệnh có thể gửi lại phải có định danh ổn định hoặc khóa chống trùng. Với kết quả chưa xác định từ MISA, JPOS hay dịch vụ khác, JPULSE lưu trạng thái để đối soát trước khi tiếp tục xử lý.", "Ranh giới dữ liệu và chống trùng: ")
add_before(anchor, "Tác vụ dài hạn như phát hành hóa đơn hàng loạt hoặc tạo voucher được tiếp nhận và theo dõi bằng trạng thái bền vững, có retry và khả năng phục hồi. API phản hồi mã công việc thay vì giữ một kết nối chờ đến khi toàn bộ tác vụ hoàn tất.", "Xử lý nền có thể phục hồi: ")
add_before(anchor, "Các enum, trạng thái và payload lõi được chia sẻ qua @bduck/shared-types. Cấu hình môi trường, bí mật và feature flag tách khỏi mã nguồn; thay đổi hợp đồng không tương thích cần có giai đoạn chuyển tiếp và được kiểm tra trong pipeline phát hành.", "Hợp đồng và cấu hình nhất quán: ")

# 3.2. Kiểu kiến trúc được lựa chọn
anchor = heading("Sơ đồ kiến trúc tổng thể")
add_before(anchor, "JPULSE được tổ chức thành ứng dụng web, dịch vụ API và các kho dữ liệu/tệp trên nền tảng Firebase. Web JPULSE dùng Next.js/React để cung cấp giao diện nghiệp vụ; JPULSE API dùng Node.js/Express để xử lý lệnh, phân quyền và tích hợp. Firestore lưu tài khoản, quyền và dữ liệu nghiệp vụ; Firebase Storage lưu chứng từ, hình ảnh và tệp. Firebase Authentication là dịch vụ xác thực danh tính bên ngoài ranh giới ứng dụng JPULSE.")
add_before(anchor, "Web trao đổi nghiệp vụ với API qua HTTP/JSON. Một số chức năng đọc dữ liệu thời gian thực và thao tác hẹp được web thực hiện qua Firebase Web SDK, dưới sự kiểm soát của Security Rules; các lệnh tạo/sửa dữ liệu nghiệp vụ đi qua API. API sử dụng Firebase Admin SDK để xác minh danh tính, truy cập Firestore/Storage và thực thi quy tắc phía máy chủ. JPOS tương tác với JPULSE qua API; các tích hợp khác được mô tả ở những góc nhìn và luồng xử lý tương ứng.")
add_before(anchor, "Ở mức C4, web, API, Firestore và Storage là các container logic: ứng dụng hoặc kho dữ liệu có trách nhiệm riêng, không đồng nghĩa với Docker container. Cách đóng gói và triển khai trên hạ tầng được trình bày tại Chương 6; Hình 3-1 tập trung vào ranh giới và những kết nối trực tiếp quan trọng của hệ thống.")

# 3.3. Sơ đồ kiến trúc tổng thể — dedicated landscape page for legibility.
diagram_heading = heading("Sơ đồ kiến trúc tổng thể")
caption = next(p for p in doc.paragraphs if p.style.name == "Caption" and p.text.strip() == "Hình 3-1: Sơ đồ kiến trúc tổng thể JPULSE")
assert DIAGRAM.exists()

portrait_end = deepcopy(doc.element.body[-1])
landscape_end = deepcopy(portrait_end)
for child in list(landscape_end):
    if child.tag in (qn("w:pgNumType"), qn("w:titlePg")):
        landscape_end.remove(child)
page_size = landscape_end.find(qn("w:pgSz"))
page_size.set(qn("w:w"), "15840")
page_size.set(qn("w:h"), "12240")
page_size.set(qn("w:orient"), "landscape")

before_break = OxmlElement("w:p")
before_props = OxmlElement("w:pPr")
before_props.append(portrait_end)
before_break.append(before_props)
diagram_heading._p.addprevious(before_break)

picture = caption.insert_paragraph_before(style="Normal")
picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
picture.paragraph_format.first_line_indent = Inches(0)
picture.paragraph_format.space_before = Inches(0)
picture.paragraph_format.space_after = Inches(0)
picture.paragraph_format.keep_with_next = True
picture.add_run().add_picture(str(DIAGRAM), width=Inches(8.8))

for run in caption.runs:
    run.text = ""
caption.add_run("Hình 3-1: Sơ đồ container (C4 cấp 2) của JPULSE")
caption.paragraph_format.keep_with_next = False

after_break = OxmlElement("w:p")
after_props = OxmlElement("w:pPr")
after_props.append(landscape_end)
after_break.append(after_props)
caption._p.addnext(after_break)

final_section = doc.element.body[-1]
for child in list(final_section):
    if child.tag in (qn("w:pgNumType"), qn("w:titlePg")):
        final_section.remove(child)

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == headings_before
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
