from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE-CapNhat-3.4.1.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-3.4-complete.docx"

doc = Document(SOURCE)
headings_before = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]


def heading(text: str):
    matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == text]
    assert len(matches) == 1, (text, len(matches))
    return matches[0]


def add_before(anchor, text: str, lead: str | None = None):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.paragraph_format.keep_together = True
    if lead:
        run = paragraph.add_run(lead)
        run.bold = True
    paragraph.add_run(text)


# 3.4.2. API nghiệp vụ JPULSE
anchor = heading("Gói kiểu dữ liệu và hợp đồng dùng chung")
add_before(anchor, "JPULSE API là ứng dụng Node.js/Express viết bằng TypeScript trong apps/be-wms. API cung cấp các tuyến REST/JSON cho xác thực, tài khoản và vai trò, cơ sở và kho, sản phẩm, nhập xuất và điều chuyển, kiểm kê, nhân sự, doanh thu, chi phí, hóa đơn, POS và voucher. Trong mã nguồn, tuyến API tiếp nhận yêu cầu và áp dụng middleware; service thực thi chính sách nghiệp vụ; repository hoặc client tích hợp truy cập dữ liệu và hệ thống ngoài. Trang yêu cầu hóa đơn công khai sử dụng tuyến và cơ chế xác thực riêng với khu vực quản trị (SRS: TC03; Chương 4 và 6).")
add_before(anchor, "API xác minh định danh hoặc phiên, xây dựng ngữ cảnh truy cập, kiểm tra quyền theo vai trò và Office Scope, kiểm tra schema và trạng thái nghiệp vụ trước khi ghi dữ liệu. Những thay đổi quan trọng như biến động tồn kho, phê duyệt, số dư nghỉ phép và phát hành hóa đơn phải bảo toàn tính nhất quán, lưu lịch sử xử lý và chống gửi trùng khi yêu cầu có thể thử lại. Middleware hiện có áp dụng Helmet, CORS theo danh sách cho phép, giới hạn tốc độ và giới hạn kích thước JSON; quyền ở giao diện không thay thế kiểm tra phía API (SRS: BC01–BC07, SEC01–SEC10).", "Kiểm soát nghiệp vụ và truy cập: ")
add_before(anchor, "API dùng Firebase Admin SDK để xác minh token và truy cập Firestore/Storage bằng quyền máy chủ, đồng thời gọi các dịch vụ bên ngoài qua client hoặc adapter riêng. Lệnh tích hợp kéo dài được giao cho tác vụ nền hoặc hàng đợi khi cấu hình cho phép; trạng thái job, khóa chống trùng và kết quả từng phần được lưu bền vững để có thể retry hoặc đối soát. Dịch vụ phải không phụ thuộc trạng thái chỉ nằm trong bộ nhớ một instance, phù hợp với cách mở rộng ngang của Cloud Run (SRS: TC04–TC05, TC11–TC13, SCA01–SCA05).", "Điều phối và vận hành: ")

# 3.4.3. Gói kiểu dữ liệu và hợp đồng dùng chung
anchor = heading("Kho dữ liệu và lưu trữ tệp")
add_before(anchor, "Gói @bduck/shared-types trong packages/shared-types là hợp đồng TypeScript được Web JPULSE và JPULSE API dùng chung. Gói tập hợp enum, trạng thái, kiểu dữ liệu và payload cốt lõi của các miền như người dùng và quyền, kho, phiếu, nhân sự, chấm công, nghỉ phép, hóa đơn, POS và voucher; một số hợp đồng có schema kiểm tra dữ liệu. Nhờ dùng cùng định nghĩa, hai ứng dụng giảm sai khác khi diễn giải mã trạng thái, trường dữ liệu và kết quả nghiệp vụ (SRS: TC01, SCA06).")
add_before(anchor, "Gói này được build và nhập như dependency trong pnpm workspace; nó không chạy thành dịch vụ độc lập và không lưu dữ liệu. Thay đổi hợp đồng cần được kiểm tra cùng frontend và backend. Khi đổi tên trường, enum hoặc ý nghĩa trạng thái theo cách không tương thích, thiết kế phải có migration hoặc giai đoạn chuyển tiếp để các phiên bản triển khai khác nhau không đọc sai dữ liệu hay phản hồi API (SRS: SCA06, COM06).", "Ranh giới và thay đổi: ")

# 3.4.4. Kho dữ liệu và lưu trữ tệp
anchor = heading("Dịch vụ tích hợp bên ngoài")
add_before(anchor, "Firestore JPULSE là kho dữ liệu NoSQL chính cho tài khoản nghiệp vụ, vai trò, quyền và phạm vi cơ sở, danh mục, chứng từ, trạng thái quy trình, dữ liệu giao dịch, cấu hình và audit. Dữ liệu được tổ chức theo miền nghiệp vụ như SRS Chương 5; các truy vấn phải giới hạn theo quyền và cơ sở, dùng index và phân trang khi tập dữ liệu lớn. Các thay đổi liên quan đồng thời đến nhiều bản ghi, đặc biệt là tồn kho hoặc số dư, cần transaction/batch và kiểm soát retry để không tạo trạng thái dở dang (SRS: TC05–TC07, AVL03–AVL04).")
add_before(anchor, "Web có thể đọc một số dữ liệu thời gian thực bằng Firebase Web SDK và thực hiện các ghi hẹp được Firestore Security Rules cho phép. Rules từ chối mặc định và kiểm tra bản ghi quyền hiện hành; các mutation nghiệp vụ quan trọng đi qua JPULSE API dùng Admin SDK để áp dụng chính sách, kiểm tra trạng thái và ghi audit. Firebase Authentication cung cấp định danh, còn trạng thái tài khoản, vai trò và quyền nghiệp vụ thuộc dữ liệu JPULSE trong Firestore.", "Đường truy cập dữ liệu: ")
add_before(anchor, "Storage JPULSE dùng Firebase Cloud Storage để lưu chứng từ, hợp đồng, hình ảnh bằng chứng và tệp nghiệp vụ. Firestore lưu metadata, định danh, trạng thái và tham chiếu tới tệp; nội dung tệp nằm trong Storage. Tùy luồng, web tải tệp bằng Storage SDK hoặc dùng URL/chính sách tải do API cấp, còn API quản lý tệp bằng Admin SDK. Thiết kế phải kiểm tra quyền sở hữu, loại và kích thước tệp, đường dẫn lưu, thời hạn URL và chính sách bucket theo môi trường; bản sao lưu phải bao gồm cả dữ liệu và tệp (SRS: TC05–TC06, SEC09, BAK01).", "Lưu trữ tệp: ")

# 3.4.5. Dịch vụ tích hợp bên ngoài
caption = next(p for p in doc.paragraphs if p.style.name == "Caption" and p.text.strip() == "Bảng 3-1: Các thành phần chính của hệ thống")
add_before(caption, "Firebase Authentication là dịch vụ xác thực danh tính và cấp token cho người dùng; JPULSE API xác minh token/phiên trước khi áp dụng quyền nghiệp vụ. Firebase Cloud Messaging chuyển thông báo đẩy đến thiết bị. Google Cloud Tasks có thể phân phối tác vụ hóa đơn hoặc voucher và gọi lại worker được bảo vệ; Cloud Scheduler được SRS nêu cho các công việc định kỳ theo cấu hình triển khai. Các dịch vụ này hỗ trợ JPULSE nhưng không nắm quy tắc nghiệp vụ hoặc quyền truy cập của người dùng (SRS: INT-01, INT-10).")
add_before(caption, "JPOS là hệ thống điểm bán độc lập. JPULSE trao đổi dữ liệu sản phẩm, thiết bị, định danh hoặc yêu cầu xác thực qua API theo từng luồng đã xác nhận. Cityfuns/JoyWorld cung cấp nguồn dữ liệu doanh thu, đơn hàng và tồn đối tác; các thao tác ghi hoặc hoàn tiền chỉ được bật theo cấu hình và quyền. SRS cũng nhắc hệ thống ERP cũ, nhưng hợp đồng tích hợp và trạng thái vận hành cụ thể cần được xác minh trước khi coi đó là luồng đang chạy (SRS: INT-04–INT-07).")
add_before(caption, "MISA meInvoice là nguồn kết quả phát hành và trạng thái hóa đơn điện tử; JPULSE chuẩn hóa dữ liệu, gửi yêu cầu, lưu định danh giao dịch và đối soát kết quả. Brevo chuyển phát email giao dịch, OTP và voucher; VietQR cung cấp dữ liệu tra cứu doanh nghiệp để hỗ trợ nhập thông tin hóa đơn. Mapbox được web sử dụng cho bản đồ hoặc tuyến điều chuyển khi có cấu hình phù hợp. Mỗi kết nối chỉ trao đổi dữ liệu cần thiết và tuân theo hợp đồng API hoặc SMTP của dịch vụ tương ứng (SRS: INT-02–INT-03, INT-08–INT-09).")
add_before(caption, "Thông tin xác thực và secret của các tích hợp máy chủ phải nằm trong cấu hình được bảo vệ, không chuyển ra bundle trình duyệt. API đặt timeout, phân loại lỗi có thể retry, lỗi kết thúc và kết quả chưa xác định; các thao tác có tác động tài chính hoặc tồn kho dùng khóa chống trùng, lưu trạng thái và hỗ trợ đối soát thủ công. Khi dịch vụ ngoài không sẵn sàng, JPULSE phải thông báo đúng trạng thái và để các nghiệp vụ độc lập tiếp tục hoạt động (SRS: TC08, TC13, AVL04–AVL08).", "Kiểm soát ranh giới tích hợp: ")

rows = [
    ("Web JPULSE", "Giao diện và trạng thái thao tác", "HTTP/JSON; Firebase Web SDK", "Next.js/React"),
    ("JPULSE API", "Quy tắc, quyền, dữ liệu và tích hợp", "REST; Firebase Admin SDK", "Node.js/Express"),
    ("Gói shared-types", "Kiểu, enum và hợp đồng dùng chung", "Được nhập bởi web và API", "Gói mã nguồn"),
    ("Firestore JPULSE", "Dữ liệu nghiệp vụ, quyền và audit", "Web SDK có phạm vi; Admin SDK", "Kho dữ liệu"),
    ("Storage JPULSE", "Chứng từ, hình ảnh và tệp", "Storage SDK; API/URL ký", "Kho tệp"),
    ("Firebase Authentication", "Xác thực danh tính và token", "Web SDK; Admin SDK", "Dịch vụ ngoài"),
    ("JPOS/Cityfuns", "Nguồn hoặc đích dữ liệu bán hàng", "API theo hợp đồng tích hợp", "Hệ thống ngoài"),
    ("Dịch vụ khác", "Hóa đơn, email, tra cứu, tác vụ nền", "API/SMTP; Cloud Tasks", "MISA, Brevo…"),
]
table = doc.tables[1]
assert [cell.text for cell in table.rows[0].cells] == ["Thành phần", "Trách nhiệm", "Giao tiếp", "Ghi chú"]
white = deepcopy(table.rows[1]._tr)
tint = deepcopy(table.rows[2]._tr)
for row in list(table.rows)[3:]:
    table._tbl.remove(row._tr)
for index in range(2, len(rows)):
    table._tbl.append(deepcopy(white if index % 2 == 0 else tint))
header_props = table.rows[0]._tr.get_or_add_trPr()
if header_props.find(qn("w:tblHeader")) is None:
    header_props.append(OxmlElement("w:tblHeader"))
for row, values in zip(table.rows[1:], rows):
    for cell, value in zip(row.cells, values):
        paragraph = cell.paragraphs[0]
        for run in paragraph.runs:
            run.text = ""
        run = paragraph.add_run(value)
        run.font.size = Pt(9.5)

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == headings_before
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
