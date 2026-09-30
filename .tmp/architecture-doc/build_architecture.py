from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import hashlib
import os
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SRS-JPULSE.docx"
OUTPUT = ROOT / "SAD-JPULSE.docx"
EXPECTED = "967e3ca379f00da00e17bf1d554420f413570b0140bf96dcbd250547b356ec48"
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == EXPECTED

source = Document(SOURCE)
doc = Document(SOURCE)
body = doc.element.body
contents_break = deepcopy(source.paragraphs[222]._p)
source_tables = [deepcopy(source.tables[i]._tbl) for i in (0, 5, 3)]

# Keep the original cover elements, cover section break, and final body section.
for element in list(body)[4:-1]:
    body.remove(element)
body.replace(body[-1], deepcopy(source.sections[2]._sectPr))

cover = body[2]
replacements = {
    "Software Requirements Specification": "Software Architecture Document",
    "Soạn bởi Mạch Lâm Quốc Hoài": "",
    " – IT và Đào tạo": "",
    "Hồ Chí Minh, 18 tháng 9 năm 2026 ": "",
}
for element in cover.iter():
    if element.tag == qn("w:t") and element.text:
        for old, new in replacements.items():
            element.text = element.text.replace(old, new)

for section in doc.sections:
    for part in (section.header, section.first_page_header, section.even_page_header):
        for element in part._element.iter():
            if element.tag == qn("w:t") and element.text:
                element.text = element.text.replace("SRS for JPULSE", "SAD for JPULSE")


def add_heading(level: int, title: str):
    return doc.add_paragraph(title, style=f"Heading {level}")


def add_caption(title: str):
    paragraph = doc.add_paragraph(title, style="Caption")
    return paragraph


def add_source_table(template_index: int, headers: list[str], caption: str):
    add_caption(caption)
    table = deepcopy(source_tables[template_index])
    rows = table.findall(qn("w:tr"))
    for row in rows[3:]:
        table.remove(row)
    for i, row in enumerate(rows[:3]):
        for j, cell in enumerate(row.findall(qn("w:tc"))):
            texts = cell.findall(".//" + qn("w:t"))
            for text in texts:
                text.text = ""
            if i == 0 and j < len(headers) and texts:
                texts[0].text = headers[j]
    body.insert(len(body) - 1, table)


doc.add_paragraph("Mục lục", style="Title")
doc.add_paragraph("[[TOC]]", style="Normal")
body.insert(len(body) - 1, contents_break)

outline = [
    (1, "Giới thiệu"),
    (2, "Mục đích tài liệu"),
    (2, "Phạm vi hệ thống"),
    (2, "Đối tượng đọc tài liệu"),
    (2, "Thuật ngữ và chữ viết tắt"),
    (2, "Tài liệu tham chiếu"),
    (1, "Bối cảnh và yêu cầu kiến trúc"),
    (2, "Bối cảnh nghiệp vụ và mục tiêu hệ thống"),
    (2, "Các bên liên quan và mối quan tâm của họ"),
    (2, "Tác nhân và hệ thống bên ngoài"),
    (2, "Ranh giới hệ thống"),
    (2, "Các chức năng ảnh hưởng đến kiến trúc"),
    (2, "Yêu cầu chất lượng"),
    (3, "Bảo mật và quyền riêng tư"),
    (3, "Hiệu năng"),
    (3, "Khả năng mở rộng"),
    (3, "Tính sẵn sàng và khả năng phục hồi"),
    (3, "Khả năng bảo trì"),
    (2, "Ràng buộc và giả định thiết kế"),
    (1, "Tổng quan và chiến lược kiến trúc"),
    (2, "Nguyên tắc thiết kế"),
    (2, "Kiểu kiến trúc được lựa chọn"),
    (2, "Sơ đồ kiến trúc tổng thể"),
    (2, "Các thành phần chính và trách nhiệm"),
    (3, "Giao diện web JPULSE"),
    (3, "API nghiệp vụ JPULSE"),
    (3, "Gói kiểu dữ liệu và hợp đồng dùng chung"),
    (3, "Kho dữ liệu và lưu trữ tệp"),
    (3, "Dịch vụ tích hợp bên ngoài"),
    (2, "Cách các thành phần giao tiếp"),
    (2, "Lý do chọn giải pháp kiến trúc"),
    (1, "Các góc nhìn kiến trúc"),
    (2, "Góc nhìn ngữ cảnh – C4 System Context"),
    (2, "Góc nhìn ứng dụng và kho dữ liệu – C4 Container"),
    (2, "Góc nhìn thành phần – C4 Component"),
    (3, "Thành phần giao diện web"),
    (3, "Thành phần API nghiệp vụ"),
    (2, "Góc nhìn động"),
    (3, "Luồng xác thực và phân quyền"),
    (3, "Luồng nhập xuất và điều chuyển kho"),
    (3, "Luồng tích hợp JPOS và hóa đơn điện tử"),
    (2, "Góc nhìn triển khai"),
    (1, "Dữ liệu, giao tiếp và luồng xử lý"),
    (2, "Tổng quan dữ liệu của hệ thống"),
    (2, "Cấu trúc dữ liệu và nơi lưu trữ"),
    (3, "Dữ liệu nghiệp vụ và phân quyền"),
    (3, "Tệp và tài liệu"),
    (3, "Dữ liệu tích hợp và nhật ký hệ thống"),
    (2, "Luồng dữ liệu giữa các thành phần"),
    (2, "Giao diện kết nối và API chính"),
    (3, "API giữa giao diện web và máy chủ"),
    (3, "Kết nối với dịch vụ bên ngoài"),
    (2, "Xác thực, phân quyền và bảo vệ dữ liệu"),
    (2, "Các luồng xử lý chính"),
    (3, "Xác thực và truy cập theo vai trò, cơ sở"),
    (3, "Quản lý nhập xuất, điều chuyển và kiểm kê"),
    (3, "Ghi nhận doanh thu và phát hành hóa đơn"),
    (3, "Chấm công và nghỉ phép"),
    (2, "Xử lý lỗi và tính nhất quán dữ liệu"),
    (1, "Triển khai và vận hành"),
    (2, "Các môi trường triển khai"),
    (3, "Môi trường phát triển"),
    (3, "Môi trường kiểm thử"),
    (3, "Môi trường vận hành"),
    (2, "Sơ đồ triển khai"),
    (2, "Cấu hình ứng dụng và quản lý thông tin bí mật"),
    (2, "Quy trình triển khai và cập nhật"),
    (2, "Ghi log và giám sát"),
    (2, "Sao lưu và khôi phục dữ liệu"),
    (2, "Khả năng mở rộng và giới hạn vận hành"),
    (1, "Quyết định kiến trúc và đánh giá"),
    (2, "Danh sách quyết định kiến trúc quan trọng"),
    (2, "Lý do lựa chọn công nghệ và phương án thiết kế"),
    (2, "Các phương án đã cân nhắc"),
    (2, "Đánh đổi của phương án được chọn"),
    (2, "Đánh giá kiến trúc theo yêu cầu chất lượng"),
    (2, "Rủi ro kỹ thuật và hướng xử lý"),
    (1, "Kết luận và phụ lục"),
    (2, "Tóm tắt kiến trúc hệ thống"),
    (2, "Những giới hạn hiện tại"),
    (2, "Hướng phát triển kiến trúc"),
    (2, "Phụ lục sơ đồ"),
    (2, "Phụ lục bảng mô tả thành phần, dữ liệu hoặc API"),
    (2, "Danh mục thuật ngữ"),
]

for level, title in outline:
    add_heading(level, title)
    if title == "Thuật ngữ và chữ viết tắt":
        add_source_table(0, ["Thuật ngữ", "Tên đầy đủ", "Ý nghĩa", "Ghi chú"], "Bảng 1-1: Thuật ngữ và chữ viết tắt")
    elif title == "Sơ đồ kiến trúc tổng thể":
        add_caption("Hình 3-1: Sơ đồ kiến trúc tổng thể JPULSE")
    elif title == "Dịch vụ tích hợp bên ngoài":
        add_source_table(0, ["Thành phần", "Trách nhiệm", "Giao tiếp", "Ghi chú"], "Bảng 3-1: Các thành phần chính của hệ thống")
    elif title == "Góc nhìn ngữ cảnh – C4 System Context":
        add_caption("Hình 4-1: Sơ đồ ngữ cảnh hệ thống JPULSE")
    elif title == "Góc nhìn ứng dụng và kho dữ liệu – C4 Container":
        add_caption("Hình 4-2: Sơ đồ ứng dụng và kho dữ liệu JPULSE")
    elif title == "Thành phần API nghiệp vụ":
        add_caption("Hình 4-3: Sơ đồ thành phần API JPULSE")
    elif title == "Luồng tích hợp JPOS và hóa đơn điện tử":
        add_caption("Hình 4-4: Sơ đồ tuần tự luồng tích hợp")
    elif title == "Góc nhìn triển khai":
        add_caption("Hình 4-5: Sơ đồ triển khai JPULSE")
    elif title == "Dữ liệu tích hợp và nhật ký hệ thống":
        add_source_table(1, ["Nhóm dữ liệu", "Nơi lưu trữ", "Ghi chú"], "Bảng 5-1: Nhóm dữ liệu và nơi lưu trữ")
    elif title == "Sơ đồ triển khai":
        add_caption("Hình 6-1: Sơ đồ triển khai JPULSE")
    elif title == "Danh sách quyết định kiến trúc quan trọng":
        add_source_table(2, ["Mã", "Bối cảnh", "Quyết định", "Hệ quả"], "Bảng 7-1: Danh sách quyết định kiến trúc")
    elif title == "Phụ lục bảng mô tả thành phần, dữ liệu hoặc API":
        add_source_table(0, ["Mã", "Tên", "Trách nhiệm", "Tham chiếu"], "Bảng 8-1: Danh mục thành phần, dữ liệu hoặc API")

doc.core_properties.title = "Tài liệu kiến trúc hệ thống JPULSE"
doc.core_properties.subject = "Khung tài liệu kiến trúc hệ thống"
doc.save(OUTPUT)

# Some first-page header parts are not exposed by the active section wrappers.
# Patch only their text, leaving images and geometry untouched.
temporary = OUTPUT.with_suffix(".tmp.docx")
with zipfile.ZipFile(OUTPUT, "r") as old, zipfile.ZipFile(temporary, "w") as new:
    for item in old.infolist():
        data = old.read(item.filename)
        if item.filename.startswith("word/header") and item.filename.endswith(".xml"):
            data = data.replace(b"SRS for JPULSE", b"SAD for JPULSE")
        new.writestr(item, data)
os.replace(temporary, OUTPUT)
print(OUTPUT)
