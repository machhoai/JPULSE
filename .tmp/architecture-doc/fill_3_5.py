from pathlib import Path

from docx import Document


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE-CapNhat-3.4.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-3.5-complete.docx"

doc = Document(SOURCE)
headings_before = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
anchors = [p for p in doc.paragraphs if p.style.name == "Heading 2" and p.text.strip() == "Lý do chọn giải pháp kiến trúc"]
assert len(anchors) == 1
anchor = anchors[0]


def add(text: str, lead: str | None = None) -> None:
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.paragraph_format.keep_together = True
    if lead:
        run = paragraph.add_run(lead)
        run.bold = True
    paragraph.add_run(text)


add(
    "Web JPULSE gửi các lệnh nghiệp vụ tới JPULSE API bằng HTTP(S) với dữ liệu JSON; phản hồi trả về kết quả, trạng thái hoặc lỗi để giao diện cập nhật thao tác. Các hợp đồng dữ liệu cốt lõi được dùng chung qua gói @bduck/shared-types, còn API là nơi kiểm tra dữ liệu đầu vào, quyền và quy tắc nghiệp vụ. Với tuyến cần đăng nhập, web chuyển thông tin xác thực do Firebase Authentication cấp; API xác minh định danh, nạp trạng thái tài khoản, vai trò, quyền và phạm vi cơ sở trước khi xử lý. Các tuyến công khai, như yêu cầu hóa đơn của khách hàng, dùng cơ chế token và giới hạn truy cập riêng theo luồng (SRS: TC01–TC03, SEC01–SEC06).",
    "Web và API: ",
)
add(
    "Web dùng Firebase Web SDK để đăng nhập. Một số màn hình đọc dữ liệu Firestore theo thời gian thực và phạm vi được phép; các ghi trực tiếp từ trình duyệt chỉ dành cho các trường hợp hẹp mà Firestore Security Rules cho phép, như trạng thái thông báo hoặc tín hiệu truyền tệp LAN. Những lệnh làm thay đổi dữ liệu nghiệp vụ đi qua API để áp dụng cùng một chính sách quyền, kiểm tra trạng thái, giao dịch và audit. API truy cập Firestore bằng Firebase Admin SDK; quyền của Admin SDK không được chuyển xuống trình duyệt (SRS: TC05–TC07, SEC07–SEC10).",
    "Xác thực và dữ liệu thời gian thực: ",
)
add(
    "Tệp chứng từ và hình ảnh được trao đổi qua Storage JPULSE. Tùy nghiệp vụ, web tải tệp bằng Storage SDK theo chính sách truy cập hoặc bằng URL/chính sách tải có thời hạn do API cấp. API dùng Admin SDK để quản lý tệp, liên kết metadata trong Firestore với đối tượng trong Storage và kiểm tra quyền trước khi cấp quyền truy cập. Khi thao tác dữ liệu và tệp không thể nằm trong cùng một giao dịch, ứng dụng phải lưu trạng thái xử lý để phát hiện, thử lại hoặc dọn dẹp tệp thiếu tham chiếu (SRS: TC05–TC06, SEC09, AVL04).",
    "Giao tiếp với kho tệp: ",
)
add(
    "JPULSE API gọi các hệ thống ngoài qua client hoặc adapter theo hợp đồng của từng bên: JPOS cho các luồng điểm bán, Cityfuns/JoyWorld cho dữ liệu bán hàng, MISA meInvoice cho hóa đơn, VietQR cho tra cứu doanh nghiệp và Brevo cho email. Web gọi Mapbox trực tiếp để hiển thị bản đồ khi tính năng được cấu hình. Mỗi kết nối chỉ truyền dữ liệu cần thiết, dùng thông tin xác thực phù hợp, đặt timeout và chuyển lỗi dịch vụ ngoài thành trạng thái nghiệp vụ có thể hiểu được. Quan hệ trao đổi với ERP cũ được SRS đề cập nhưng chiều gọi và hợp đồng triển khai cần xác minh (SRS: INT-02–INT-09).",
    "Giao tiếp với hệ thống ngoài: ",
)
add(
    "API có thể tạo Cloud Task cho hóa đơn hoặc voucher khi hàng đợi được cấu hình. Worker xác thực lời gọi, chống xử lý trùng và ghi kết quả hoặc lỗi để đối soát. API gửi thông báo đẩy qua Firebase Cloud Messaging; web đọc trạng thái tác vụ từ dữ liệu JPULSE (SRS: INT-01, INT-10, TC11–TC13).",
    "Tác vụ nền và thông báo: ",
)

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == headings_before
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
