from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE-CapNhat-3.6.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-Chuong-4-complete.docx"
ARCH = ROOT / "docs" / "architecture"

doc = Document(SOURCE)
headings_before = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
old_captions = {
    "Hình 4-1: Sơ đồ ngữ cảnh hệ thống JPULSE",
    "Hình 4-2: Sơ đồ ứng dụng và kho dữ liệu JPULSE",
    "Hình 4-3: Sơ đồ thành phần API JPULSE",
    "Hình 4-4: Sơ đồ tuần tự luồng tích hợp",
    "Hình 4-5: Sơ đồ triển khai JPULSE",
}
for paragraph in list(doc.paragraphs):
    if paragraph.style.name == "Caption" and paragraph.text.strip() in old_captions:
        paragraph._p.getparent().remove(paragraph._p)


def heading(text: str):
    found = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == text]
    assert len(found) == 1, (text, len(found))
    return found[0]


def prose(anchor, text: str, lead: str | None = None):
    p = anchor.insert_paragraph_before(style="Normal")
    p.paragraph_format.keep_together = True
    if lead:
        r = p.add_run(lead)
        r.bold = True
    p.add_run(text)
    return p


figures = []


def figure(anchor, number: int, title: str, filename: str):
    path = ARCH / filename
    assert path.is_file(), path
    caption = anchor.insert_paragraph_before(style="Caption")
    caption.add_run(f"Hình 4-{number}: {title}")
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.keep_together = True
    p = anchor.insert_paragraph_before(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    with Image.open(path) as img:
        width_px, height_px = img.size
    max_width, max_height = 6.45, 6.25
    scale = min(max_width / width_px, max_height / height_px)
    width_inches = width_px * scale
    shape = p.add_run().add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", title)
    figures.append((number, title, filename, round(width_inches, 2)))


# 4.1. C4 System Context
anchor = heading("Góc nhìn ứng dụng và kho dữ liệu – C4 Container")
prose(anchor, "Góc nhìn ngữ cảnh xem JPULSE là một hệ thống phần mềm duy nhất trong hoạt động Joy World. Ban quản lý và quản trị dùng hệ thống để giám sát, phê duyệt và cấu hình; nhân viên vận hành xử lý nghiệp vụ theo quyền và cơ sở; khách hàng chỉ gửi yêu cầu hóa đơn trong luồng công khai. Hình 4-1 cho thấy những tác nhân này cùng các hệ thống do bên khác vận hành, không phân rã thành phần nội bộ của JPULSE (SRS, Mục 2.3–2.4).")
figure(anchor, 1, "Sơ đồ ngữ cảnh hệ thống JPULSE", "jpulse-system-context.png")
prose(anchor, "JPULSE trao đổi danh mục và thông tin liên quan điểm bán với JPOS; lấy dữ liệu bán hàng từ JoyWorld/Cityfuns; dùng Firebase Authentication cho định danh, MISA meInvoice cho hóa đơn, VietQR cho tra cứu doanh nghiệp, Brevo cho email, Mapbox cho bản đồ và Firebase Cloud Messaging cho thông báo đẩy. Mũi tên trên hình chỉ bên khởi tạo tương tác; việc có một đường tích hợp trong mã không khẳng định nó đã được bật ở mọi môi trường. ERP cũ không có đường trao đổi kỹ thuật đã được xác minh nên không xuất hiện trên hình. Các trách nhiệm và ranh giới chi tiết được nêu ở Mục 2.3–2.4.")


# 4.2. C4 Container
anchor = heading("Góc nhìn thành phần – C4 Component")
prose(anchor, "Hình 4-2 phân rã JPULSE thành hai ứng dụng và hai kho dữ liệu logic. Web JPULSE cung cấp giao diện; JPULSE API thực thi lệnh nghiệp vụ, quyền và tích hợp; Firestore lưu dữ liệu có cấu trúc; Storage lưu tệp. Firebase Authentication là dịch vụ xác thực bên ngoài ứng dụng. Trong mô hình C4, các khối này là container theo trách nhiệm chạy và lưu trữ, không đồng nghĩa với Docker container.")
figure(anchor, 2, "Sơ đồ container của JPULSE", "jpulse-container-main.png")
prose(anchor, "Đường chính cho thao tác tạo và sửa là Web → API → Firestore/Storage, nơi API kiểm tra quyền và trạng thái trước khi ghi. Web cũng dùng Firebase Web SDK để đăng nhập, đọc một số dữ liệu thời gian thực hoặc tải tệp trong phạm vi Security Rules cho phép. API dùng Admin SDK phía máy chủ. Gói @bduck/shared-types là mã hợp đồng được hai ứng dụng nhập khi build, nên không là container chạy độc lập. Hình 4-2 cũng đã được dùng ở Mục 3.3 để trình bày kiến trúc tổng thể.")
figure(anchor, 3, "Sơ đồ tích hợp với hệ thống và dịch vụ bên ngoài", "jpulse-integrations.png")
prose(anchor, "Hình 4-3 bổ sung các đường tích hợp không thể hiện hết trong sơ đồ container chính: API gọi hoặc nhận lời gọi từ JPOS, Cityfuns/JoyWorld, MISA, VietQR, Brevo, FCM và Cloud Tasks; Web gọi Mapbox. Đường nét đứt với ERP cũ chỉ phản ánh quan hệ SRS đề cập, còn hợp đồng và trạng thái triển khai chưa được xác minh. API phải giữ bí mật tích hợp ở phía máy chủ, kiểm soát timeout và lưu trạng thái đối soát cho các lời gọi có thể có kết quả chưa xác định (SRS: INT-01–INT-10).")


# 4.3. C4 Component
anchor = heading("Thành phần giao diện web")
prose(anchor, "Góc nhìn thành phần chỉ phân rã hai ứng dụng có cấu trúc bên trong cần giải thích. Các mô tả dưới đây nêu trách nhiệm và chiều phụ thuộc từ mã nguồn; kho biểu đồ hiện chưa có C4 Component Diagram riêng cho Web hoặc API. Vì vậy, không suy diễn các module thành những tiến trình triển khai độc lập.")

anchor = heading("Thành phần API nghiệp vụ")
prose(anchor, "Trong Web JPULSE, lớp route và layout của Next.js định tuyến các khu vực dashboard, nghiệp vụ và trang yêu cầu hóa đơn công khai. Các feature component nhận dữ liệu, hiển thị trạng thái và thu thập thao tác; hook/store giữ trạng thái phiên, quyền, bộ lọc và dữ liệu hiển thị. Lớp API client gửi lệnh đến JPULSE API, còn thư viện Firebase phục vụ xác thực, listener Firestore và một số luồng tệp được cho phép. Các lớp này dùng hợp đồng @bduck/shared-types để diễn giải payload và trạng thái nhất quán (SRS: TC01–TC02, SEC01–SEC03).")
prose(anchor, "Luồng phụ thuộc thông thường là route/component → hook hoặc client → JPULSE API hay Firebase SDK. Kiểm tra tại giao diện giúp người dùng biết thao tác hợp lệ và phản hồi lỗi sớm; nó không thay thế kiểm tra quyền, phạm vi cơ sở và trạng thái nghiệp vụ tại API. Ranh giới giữa truy cập Firebase trực tiếp và lệnh phải qua API cần được giữ rõ khi thêm màn hình mới.")

anchor = heading("Góc nhìn động")
prose(anchor, "Trong JPULSE API, route tiếp nhận HTTP request và gắn middleware xác thực, giới hạn truy cập; controller chuẩn hóa đầu vào và phản hồi; service giữ quy tắc, trạng thái và điều phối giao dịch; repository hoặc Firebase Admin SDK đọc ghi Firestore/Storage. Client tích hợp bao bọc hợp đồng JPOS, MISA và các hệ thống ngoài; worker xử lý tác vụ được giao lại từ hàng đợi. Cách phân lớp này giữ quy tắc nghiệp vụ ở máy chủ, đồng thời cô lập thay đổi giao thức bên ngoài (SRS: TC03–TC05, TC11–TC13).")
prose(anchor, "Chiều gọi chính là route/middleware → controller → service → repository hoặc client tích hợp. Middleware và service đều cần ngữ cảnh định danh, vai trò và Office Scope; service phải kiểm tra lại điều kiện nghiệp vụ ngay trước thay đổi dữ liệu. Các thao tác nhiều bước lưu trạng thái và audit để nhận biết phần đã thành công khi bước sau lỗi. C4 Component Diagram cho API vẫn cần được vẽ để thể hiện chính thức các quan hệ này; phần mô tả hiện tại là cơ sở cho hình đó.")


# 4.4. Dynamic view
anchor = heading("Luồng xác thực và phân quyền")
prose(anchor, "Góc nhìn động chọn những sequence diagram thể hiện ranh giới xác thực, ghi dữ liệu kho và tương tác với hệ thống ngoài. Mỗi hình tập trung một bước có giá trị kiến trúc; những bước kề trước/sau được chỉ rõ bằng mã SD trong phần mô tả. Bộ hình đầy đủ nằm trong docs/architecture/sequence/. Mũi tên liền biểu diễn lời gọi, mũi tên đứt biểu diễn phản hồi hoặc dữ liệu đẩy.")

anchor = heading("Luồng nhập xuất và điều chuyển kho")
prose(anchor, "Luồng đăng nhập Web bắt đầu bằng phân giải định danh và xác thực với Firebase Authentication ở SD-02A. Hình 4-4 tiếp tục từ ID token: API xác minh token, đọc tài khoản và quyền JPULSE trong Firestore, từ chối tài khoản không hợp lệ rồi mới tạo phiên. Firebase xác nhận danh tính; quyền nghiệp vụ và phạm vi cơ sở do JPULSE quyết định.")
figure(anchor, 4, "SD-02B tạo phiên và nạp quyền JPULSE", "sequence/sd-02b-jpulse-session.png")
prose(anchor, "Nếu tài khoản yêu cầu MFA, Web giữ màn hình khóa sau khi tạo phiên. Hình 4-5 cho thấy API xác minh mã TOTP hoặc email OTP; mã sai/hết hạn không mở thao tác, còn mã hợp lệ cho phép Web tiếp tục. Trạng thái mở khóa ở sơ đồ là trạng thái giao diện sau khi API xác minh, không làm thay đổi nguyên tắc kiểm tra quyền tại API cho mỗi lệnh (SRS: SEC01–SEC03).")
figure(anchor, 5, "SD-02C xác minh MFA sau đăng nhập", "sequence/sd-02c-jpulse-mfa.png")

anchor = heading("Luồng tích hợp JPOS và hóa đơn điện tử")
prose(anchor, "Với phiếu nhập, các bước SD-03A/B tải chứng từ và kiểm tra quyền, dữ liệu trước khi ghi. Hình 4-6 thể hiện API ghi phiếu, dòng hàng và bước duyệt bằng batch Firestore, sau đó ghi audit và trả trạng thái. Batch thất bại không tạo phiếu; audit được ghi sau batch nên lỗi ở bước audit cần được giám sát và đối soát riêng, không được mô tả là toàn bộ thao tác đã tự rollback.")
figure(anchor, 6, "SD-03C ghi phiếu nhập và audit", "sequence/sd-03c-import-persist.png")
prose(anchor, "Phiếu xuất tuân theo cùng ranh giới Web → API → Firestore. Sau khi SD-06A kiểm tra chứng từ, quyền, vị trí và OTP, Hình 4-7 mô tả bước ghi phiếu và khởi tạo duyệt. Một số bước audit hoặc hoàn tất khởi tạo duyệt chạy sau batch; nếu chúng lỗi, dữ liệu phiếu có thể đã tồn tại. Luồng xuất tiếp tục ở các sơ đồ duyệt, soạn hàng và chốt phiếu SD-07–09 (SRS: BC01–BC05).")
figure(anchor, 7, "SD-06B ghi phiếu xuất và khởi tạo duyệt", "sequence/sd-06b-export-persist.png")
prose(anchor, "Điều chuyển liên kho cần giữ quan hệ giữa phiếu xuất ở kho nguồn và việc nhận ở kho đích. Hình 4-8 mô tả bước nhận: kho đích ghi số lượng thực nhận; transaction giảm lượng đang chuyển ở nguồn, tăng ATP tại đích và lưu chênh lệch từng dòng. Chênh lệch nhận trong luồng này không tự tạo báo cáo không phù hợp; quy trình xử lý chênh lệch cần được theo dõi riêng (SRS: BC04–BC05).")
figure(anchor, 8, "SD-17 nhận hàng điều chuyển", "sequence/sd-17-transfer-receiving.png")

anchor = heading("Góc nhìn triển khai")
prose(anchor, "Hai loại tương tác cần được phân biệt. Hình 4-9 mô tả hiện trạng JPOS nạp quyền: sau khi Firebase xác thực, JPOS Functions đọc tài khoản, vai trò và phạm vi cơ sở từ Firestore dùng chung rồi tính quyền pos.login. Cách này khác mô tả mục tiêu trong SRS/Container Diagram rằng JPOS gọi JPULSE API để xác thực. Hình ghi lại hành vi đã được đối chiếu với mã JPOS; hợp đồng tích hợp API cho xác thực hoặc đồng bộ cần được xác minh trước khi coi hai mô tả là một. Các sơ đồ SD-01A/C lần lượt nêu bước xác thực ban đầu và giới hạn cache khi mất mạng.")
figure(anchor, 9, "SD-01B JPOS nạp quyền từ dữ liệu JPULSE", "sequence/sd-01b-jpos-access.png")
prose(anchor, "Với hóa đơn điện tử, Hình 4-10 cho thấy API kiểm tra phiên, quyền, cấu hình và bản nháp trước khi ghi job/item và tạo Cloud Task. Web nhận mã công việc và trạng thái hàng đợi mà không chờ MISA. Khóa tham chiếu và tên task ổn định giúp ngăn tạo trùng khi yêu cầu được gửi lại (SRS: P04, INT-03, INT-10).")
figure(anchor, 10, "SD-04A tiếp nhận và xếp hàng phát hành hóa đơn", "sequence/sd-04a-invoice-queue.png")
prose(anchor, "Hình 4-11 mô tả Cloud Tasks gọi worker của JPULSE API. Worker xác thực lời gọi, claim item/lane trong Firestore và chỉ gửi yêu cầu phát hành tới MISA khi claim thành công. Task đến sớm, lane bận hoặc item đã kết thúc được xử lý theo nhánh riêng, tránh phát hành lại ngoài ý muốn.")
figure(anchor, 11, "SD-04B worker nhận task và gọi MISA", "sequence/sd-04b-invoice-worker.png")
prose(anchor, "Hình 4-12 thể hiện bước ghi kết quả MISA: thành công thành ISSUED; timeout hoặc phản hồi mơ hồ thành PENDING_CONFIRMATION; lỗi xác định được phân loại. Kết quả chưa rõ không được tự phát hành lại ngay mà phải qua bước xác nhận/đối soát ở SD-04D. Cách lưu trạng thái này cho phép người dùng theo dõi và xử lý thủ công mà không nhân đôi hóa đơn (SRS: BC06–BC07, AVL04–AVL08).")
figure(anchor, 12, "SD-04C ghi kết quả phát hành hóa đơn", "sequence/sd-04c-invoice-outcome.png")


# 4.5. Deployment view
anchor = heading("Dữ liệu, giao tiếp và luồng xử lý")
prose(anchor, "Ở mức triển khai, SRS xác định backend Node.js/Express được đóng gói Docker và định hướng chạy trên Google Cloud Run tại asia-southeast1; Web Next.js được build ở chế độ standalone. Firestore, Cloud Storage và Firebase Authentication là dịch vụ được quản lý; Cloud Tasks tham gia các luồng nền khi queue, URL worker và quyền gọi được cấu hình. Các thành phần phải dùng cấu hình, secret và service account theo môi trường; API/worker cần xác thực lời gọi, ghi log và health check. Trạng thái nghiệp vụ nằm trong kho bền vững để instance Cloud Run có thể thay đổi hoặc mở rộng (SRS: TC01–TC06, TC08, TC11–TC13, SCA01–SCA02).")
prose(anchor, "Kho biểu đồ chưa có Deployment Diagram xác định nơi chạy Web, ingress, IAM, queue/worker và vị trí Firestore/Storage theo môi trường. Vì vậy Mục 4.5 chỉ ghi các ràng buộc đã được SRS xác nhận, chưa khẳng định topology production. Cần bổ sung hình sau khi chốt môi trường thực tế và đối chiếu kế hoạch vận hành ở Chương 6.")

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == headings_before
assert [item[0] for item in figures] == list(range(1, 13))
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
print("FIGURES", figures)
