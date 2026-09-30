from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE-Chuong-1.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-chapter-2.docx"

doc = Document(SOURCE)
original_headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]


def heading(text: str):
    matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == text]
    assert len(matches) == 1, (text, len(matches))
    return matches[0]


def add_before(anchor, text: str, lead: str | None = None):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    if lead:
        run = paragraph.add_run(lead)
        run.bold = True
    paragraph.add_run(text)


# 2.1. Bối cảnh nghiệp vụ và mục tiêu hệ thống
anchor = heading("Các bên liên quan và mối quan tâm của họ")
add_before(anchor, "Joy World từng quản lý nhiều hoạt động cửa hàng bằng hồ sơ giấy, bảng tính và hệ thống có phạm vi hạn chế. Theo SRS, dữ liệu phân tán, đối soát thủ công và mức độ liên kết cao trong phần mềm cũ gây khó khăn cho việc kiểm soát, bảo trì và mở rộng. JPULSE được hình thành từ JWC-ERPS để thay thế dần cách vận hành đó bằng một nền tảng quản trị có giao diện và dịch vụ xử lý nghiệp vụ tách biệt.")
add_before(anchor, "Mục tiêu của JPULSE là số hóa quy trình, tập trung dữ liệu, bảo toàn lịch sử xử lý và giúp các bộ phận nhìn thấy trạng thái nghiệp vụ theo quyền được cấp. Hệ thống khởi đầu từ quản lý kho, sau đó mở rộng sang nhân sự, chi phí, doanh thu, hóa đơn điện tử và các nghiệp vụ quản trị khác. Kết nối với JPOS giúp dữ liệu bán hàng đến hệ thống quản trị gần thời gian thực; kết nối MISA meInvoice giảm các bước tổng hợp và phát hành hóa đơn thủ công.")
add_before(anchor, "Những mục tiêu trên dẫn đến các nhu cầu kiến trúc chính: tách trách nhiệm giữa giao diện, API và dữ liệu; duy trì một hợp đồng trao đổi nhất quán; kiểm soát quyền theo vai trò và cơ sở; xử lý an toàn các tác vụ kéo dài hoặc kết quả tích hợp chưa xác định; và mở rộng theo số cơ sở, người dùng và giao dịch mà không làm mất tính đúng đắn của dữ liệu (SRS, Chương 2).")

# 2.2. Các bên liên quan và mối quan tâm của họ
anchor = heading("Tác nhân và hệ thống bên ngoài")
add_before(anchor, "Ban Giám đốc và Quản lý Vận hành cần số liệu kịp thời, quy trình thống nhất và khả năng giám sát theo cơ sở. Vì vậy, kiến trúc phải cung cấp dữ liệu báo cáo có nguồn gốc rõ ràng, cập nhật có kiểm soát và chỉ hiển thị trong phạm vi được phép.")
add_before(anchor, "Kho vận, Kế toán, Nhân sự và Marketing là các đơn vị sở hữu quy tắc nghiệp vụ. Kho vận quan tâm đến số dư tồn kho và chứng từ có thể truy vết; Kế toán quan tâm đến đối soát doanh thu, chi phí và hóa đơn; Nhân sự cần bảo vệ hồ sơ, chấm công và nghỉ phép; Marketing cần kiểm soát vòng đời chiến dịch và mã voucher. Quản lý cơ sở, nhân viên kho, cửa hàng và thu ngân cần thao tác rõ ràng trên thiết bị phù hợp, kể cả các luồng được phép hoạt động khi kết nối không ổn định.")
add_before(anchor, "Quản trị hệ thống và IT cần quản lý tài khoản, quyền, cấu hình, bí mật, giám sát và khôi phục. Kiểm soát nội bộ cần audit log và lịch sử phê duyệt để xác định ai thực hiện thay đổi, khi nào và trong phạm vi nào. Nhóm Product, BA, Dev và QA cần ranh giới thành phần, hợp đồng dữ liệu và tiêu chí chất lượng đủ rõ để phát triển, kiểm thử và đánh giá tác động thay đổi (SRS, Mục 2.3).")

# 2.3. Tác nhân và hệ thống bên ngoài
anchor = heading("Ranh giới hệ thống")
add_before(anchor, "Tác nhân sử dụng JPULSE gồm Ban Giám đốc, Quản lý Vận hành, các phòng ban Kho vận, Kế toán, Nhân sự, Marketing, quản lý cơ sở, nhân viên cửa hàng, thu ngân, quản trị viên và người được cấp quyền kiểm toán. Khách hàng chỉ tương tác trong phạm vi cung cấp thông tin và nhận hóa đơn hoặc thông báo liên quan đến giao dịch; họ không truy cập dữ liệu quản trị nội bộ.")
add_before(anchor, "Các hệ thống trao đổi dữ liệu với JPULSE gồm JPOS/JoyWorld và hệ thống ERP cũ cho dữ liệu bán hàng, người dùng hoặc tồn kho; MISA meInvoice cho kết quả phát hành hóa đơn; Brevo cho email giao dịch; VietQR cho tra cứu thông tin doanh nghiệp. Google Cloud và Firebase cung cấp môi trường chạy, xác thực, Firestore và Cloud Storage. Cơ quan quản lý nhà nước đặt ra yêu cầu tuân thủ, không phải một điểm kết nối kỹ thuật mặc định (SRS, Mục 2.3–2.4).")
add_before(anchor, "Mỗi kết nối phải có chủ thể sở hữu dữ liệu và trách nhiệm xử lý rõ ràng. Hệ thống nguồn giữ quyền đối với dữ liệu gốc do mình phát hành; JPULSE xác thực yêu cầu, chuẩn hóa định danh, lưu trạng thái xử lý và đối soát kết quả cần thiết cho quy trình nội bộ. Kết quả trả về chưa xác định từ hệ thống ngoài không được coi là thành công.")

# 2.4. Ranh giới hệ thống
anchor = heading("Các chức năng ảnh hưởng đến kiến trúc")
add_before(anchor, "Trong ranh giới JPULSE có giao diện người dùng do dự án cung cấp, API và tác vụ nền, quy tắc nghiệp vụ, dữ liệu nội bộ, tệp chứng từ, báo cáo, nhật ký kiểm toán và lớp điều phối tích hợp. JPULSE chịu trách nhiệm kiểm tra định danh, quyền theo vai trò và phạm vi cơ sở, hợp lệ hóa đầu vào, chuyển trạng thái quy trình, lưu thay đổi và cung cấp kết quả cho người dùng hoặc hệ thống được phép.")
add_before(anchor, "Người dùng và các hệ thống kể ở Mục 2.3 nằm ngoài ranh giới triển khai của JPULSE. Mọi trao đổi phải đi qua giao diện, API, webhook hoặc tác vụ tích hợp đã xác định; hệ thống ngoài không truy cập trực tiếp kho dữ liệu nội bộ. JPULSE không kiểm soát hoạt động bên trong, thời gian sẵn sàng hay dữ liệu có thẩm quyền của MISA, Brevo, VietQR và các hệ thống bán hàng.")
add_before(anchor, "Theo SRS, JPULSE không thực hiện tính lương, kế toán tổng hợp hoặc cổng thương mại điện tử cho khách hàng. Ranh giới này giúp các quyết định lưu trữ, phân quyền và tích hợp tập trung vào dữ liệu và quy trình mà JPULSE thực sự sở hữu (SRS, Mục 1.2 và 2.4).")

# 2.5. Các chức năng ảnh hưởng đến kiến trúc
anchor = heading("Yêu cầu chất lượng")
add_before(anchor, "Nghiệp vụ kho tạo ra các thay đổi liên quan đồng thời đến phiếu, phê duyệt và nhiều bucket tồn kho. Kiến trúc phải bảo đảm tính nguyên tử, kiểm tra lượng khả dụng, ngăn thao tác vượt quyền và lưu vết đủ để đối soát khi kiểm kê hoặc điều chuyển. Dữ liệu nhập tại hiện trường có thể cần đồng bộ lại sau khi mất kết nối (SRS: BC01–BC05, BC14).")
add_before(anchor, "Các luồng hóa đơn, doanh thu và POS chịu ảnh hưởng của hệ thống ngoài và trạng thái giao dịch bất đồng bộ. JPULSE phải ngăn phát hành trùng hóa đơn, nhận diện trạng thái chưa xác định, lưu khóa chống trùng và hỗ trợ đối soát. Dữ liệu JPOS được tổng hợp theo nguồn, trạng thái thanh toán và phạm vi cơ sở; tác vụ phát hành hoặc đồng bộ lớn cần chạy nền và hiển thị tiến độ (SRS: BC06–BC07, P04).")
add_before(anchor, "Hồ sơ nhân viên, chấm công, nghỉ phép, chi phí và voucher dùng chính sách, phê duyệt, số dư hoặc feature flag khác nhau. Chúng cần một cơ chế nhất quán để kiểm tra quyền, trạng thái, thời điểm nghiệp vụ và lịch sử thay đổi, nhưng vẫn giữ quy tắc riêng ở từng miền nghiệp vụ. Các thao tác nhạy cảm như xử lý chênh lệch tồn kho hoặc phát hành hóa đơn cần xác thực tăng cường theo SRS.")
add_before(anchor, "Tất cả luồng dùng chung định danh người dùng, vai trò, phạm vi cơ sở, cấu hình tích hợp, thông báo và audit. Vì vậy, kiểm soát truy cập và hợp đồng dữ liệu giữa giao diện, API, tác vụ nền và hệ thống ngoài là các mối quan tâm xuyên suốt, không thể chỉ kiểm tra tại màn hình người dùng.")

# 2.6. Yêu cầu chất lượng
anchor = heading("Bảo mật và quyền riêng tư")
add_before(anchor, "Các mục dưới đây tóm tắt những yêu cầu có tác động trực tiếp đến kiến trúc. Giá trị mục tiêu và tiêu chí nghiệm thu chi tiết thuộc Chương 7 của SRS; khi triển khai cần đo và đối chiếu theo môi trường tương ứng.")

anchor = heading("Hiệu năng")
add_before(anchor, "Người dùng web xác thực qua Firebase và phiên đăng nhập an toàn; API phải kiểm tra quyền theo vai trò và phạm vi cơ sở trước khi truy cập hoặc thay đổi dữ liệu. Khi ngữ cảnh quyền không tồn tại hoặc không hợp lệ, hệ thống từ chối mặc định. Firestore Rules áp dụng deny-by-default; các thay đổi nghiệp vụ quan trọng đi qua backend để kiểm tra trạng thái và ghi audit (SRS: SEC01–SEC03, TC06–TC07).")
add_before(anchor, "Bí mật tích hợp, khóa mã hóa và thông tin service account chỉ được cấp cho môi trường máy chủ; kết nối production dùng HTTPS/TLS, CORS theo danh sách cho phép và giới hạn tốc độ truy cập. Dữ liệu cá nhân, thông tin nhân sự, OTP, token và nội dung hóa đơn phải được tối thiểu hóa trong trao đổi, giới hạn quyền đọc và không ghi nguyên văn vào log. Thao tác rủi ro cao dùng OTP/MFA theo chính sách (SRS: SEC04–SEC10).")

anchor = heading("Khả năng mở rộng")
add_before(anchor, "SRS đặt mục tiêu 95% yêu cầu API thông thường không gọi hệ thống ngoài hoàn tất trong 2 giây và 99% trong 5 giây ở tải danh định; màn hình nghiệp vụ đầu tiên có thể thao tác trong 3 giây trên mạng văn phòng ổn định. Danh sách dữ liệu dài cần phân trang hoặc tải tăng dần để giới hạn lượng dữ liệu truyền và xử lý trên client (SRS: P01–P02).")
add_before(anchor, "Lệnh phát hành hóa đơn hàng loạt, xuất voucher hoặc đồng bộ lớn cần được tiếp nhận nhanh, trả mã công việc và xử lý nền thay vì giữ kết nối chờ. Mọi lời gọi hệ thống ngoài phải có timeout; SRS nêu ngưỡng 15 giây cho MISA và yêu cầu phân biệt lỗi có thể thử lại với trạng thái chưa xác định (SRS: P03–P04).")

anchor = heading("Tính sẵn sàng và khả năng phục hồi")
add_before(anchor, "Backend phải không phụ thuộc trạng thái chỉ nằm trong bộ nhớ hoặc ổ đĩa cục bộ để có thể mở rộng ngang. SRS ghi nhận cấu hình Cloud Run hiện tại tại asia-southeast1 với min 0, max 3 instance và concurrency 80; việc tăng giới hạn phải qua kiểm thử tải mà không cần thay đổi tính đúng đắn của nghiệp vụ (SRS: SCA01–SCA02, TC04).")
add_before(anchor, "Truy vấn và báo cáo phải giới hạn theo cơ sở, thời gian và số bản ghi, dùng index cùng cơ chế phân trang phù hợp Firestore. Các tác vụ fan-out đặt giới hạn song song; tác vụ nền dùng khóa chống trùng, lease và retry để nhiều instance không xử lý cùng một đơn vị công việc. Hợp đồng dữ liệu dùng chung cần có phương án chuyển tiếp khi thay đổi không tương thích (SRS: SCA03–SCA06).")

anchor = heading("Khả năng bảo trì")
add_before(anchor, "SRS đặt mục tiêu mức sẵn sàng tối thiểu 99,5% theo tháng cho API lõi. Kiến trúc cần health check, trạng thái công việc bền vững và cơ chế suy giảm có kiểm soát khi dịch vụ tích hợp không sẵn sàng; các nghiệp vụ độc lập vẫn phải hoạt động. Cập nhật tồn kho, số dư nghỉ phép và khóa công việc cần transaction hoặc batch; yêu cầu có thể gửi lại phải có định danh ổn định để không nhân đôi tác động (SRS: AVL01–AVL05, AVL08).")
add_before(anchor, "Các luồng được phép hoạt động ngoại tuyến phải giữ thời điểm thao tác, đánh dấu dữ liệu chưa đồng bộ và chống gửi trùng khi có mạng trở lại. Tác vụ nền có checkpoint hoặc lease để phục hồi sau lỗi. SRS đặt mục tiêu khôi phục dữ liệu giao dịch lõi với RPO không quá 1 giờ và RTO không quá 4 giờ; lịch sao lưu và thử phục hồi phải được vận hành, kiểm chứng riêng (SRS: AVL06–AVL07, BAK01–BA04).")

anchor = heading("Ràng buộc và giả định thiết kế")
add_before(anchor, "Khả năng bảo trì dựa trên việc tách giao diện khỏi API, giữ logic nghiệp vụ quan trọng phía máy chủ và dùng @bduck/shared-types cho enum, trạng thái và payload lõi. Thay đổi hợp đồng không tương thích cần version hoặc giai đoạn chuyển tiếp; thay đổi dữ liệu cần migration/backfill có dry-run, checkpoint và phương án hoàn nguyên (SRS: TC01–TC03, SCA06, BAK06).")
add_before(anchor, "Cấu hình, feature flag và bí mật phải được quản lý theo môi trường. Pipeline phát hành cần chạy lint, typecheck và các kiểm thử phân quyền, danh tính, Firestore Rules cùng tích hợp liên quan trước khi build hoặc triển khai; đây là cổng kiểm soát đặc biệt quan trọng vì SRS ghi nhận cấu hình build frontend hiện bỏ qua lỗi lint và TypeScript (SRS: TC08, TC10, TC16, SEC11).")

# 2.7. Ràng buộc và giả định thiết kế
anchor = heading("Tổng quan và chiến lược kiến trúc")
add_before(anchor, "Ràng buộc nền tảng từ SRS gồm monorepo pnpm 9.15.4, Node.js từ phiên bản 22 và TypeScript; frontend Next.js/React ở chế độ standalone, backend Node.js/Express cung cấp REST API và gói @bduck/shared-types làm hợp đồng dùng chung. Thiết kế phải dựa trên API và quy ước của phiên bản framework thực tế trong mã nguồn, không giả định khả năng của phiên bản khác (SRS: TC01–TC03).")
add_before(anchor, "Firestore là kho dữ liệu nghiệp vụ chính, Firebase Authentication phục vụ định danh và Cloud Storage lưu tệp/chứng từ. Backend được đóng gói Docker và SRS định hướng triển khai trên Google Cloud Run tại asia-southeast1. Thiết kế phải tính đến truy vấn theo cơ sở, index, transaction và Firestore Rules (SRS: TC04–TC06).")
add_before(anchor, "Các luồng phụ thuộc MISA, JPOS/JoyWorld, Brevo, VietQR và Open API cần hợp đồng trao đổi, xác thực, timeout, idempotency và đối soát riêng. Tác vụ hóa đơn/voucher dài hạn cần hàng đợi hoặc worker được bảo vệ. Thời điểm nghiệp vụ và đồng bộ thống nhất theo múi giờ Asia/Ho_Chi_Minh (SRS: TC11–TC13).")
add_before(anchor, "Thiết kế giả định đơn vị vận hành cấp tài khoản dịch vụ, bí mật, IAM, mạng và môi trường Firebase/GCP; các phòng ban xác nhận quy tắc, phạm vi cơ sở và ánh xạ định danh. Vì chất lượng và tính sẵn sàng của dịch vụ ngoài không thuộc JPULSE, các luồng phụ thuộc cần trạng thái lỗi, retry và đối soát. Các giả định này phải được xác nhận trước triển khai.")

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == original_headings
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
