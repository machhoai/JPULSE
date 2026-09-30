import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(r"D:\Github\bduck-system")
ARCH = ROOT / "docs" / "architecture"
SEQUENCE = ARCH / "sequence"
SOURCE = ROOT / "SAD-JPULSE-CapNhat-Chuong-4.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-Chuong-4-all-sd.docx"
MD_FILES = [
    "jpulse-sequence-diagrams.md",
    "jpulse-warehouse-sequence-diagrams.md",
    "jpulse-facility-sequence-diagrams.md",
    "jpulse-tasks-stock-count-sequence-diagrams.md",
    "jpulse-invoice-management-sequence-diagrams.md",
    "jpulse-marketing-voucher-sequence-diagrams.md",
    "jpulse-expense-sequence-diagrams.md",
]


def diagram_titles():
    titles = {}
    image_pattern = re.compile(r"!\[([^]]+)\]\(\./sequence/([^)]+\.png)\)")
    for name in MD_FILES:
        heading = ""
        for line in (ARCH / name).read_text(encoding="utf-8").splitlines():
            if line.startswith("#") and "SD-" in line:
                heading = re.sub(r"^#+\s*", "", line)
                heading = re.sub(r"^\d+\.\s*", "", heading)
            match = image_pattern.search(line)
            if not match:
                continue
            alt, filename = match.groups()
            stem = Path(filename).stem
            code = re.match(r"(sd-\d+[a-z]?)", stem).group(1).upper()
            alt_title = re.sub(r"^SD-\d+[A-Z]?\s*[-–—:]?\s*", "", alt, flags=re.I).strip()
            head_title = re.sub(r"^SD-[^–—]+\s*[–—]\s*", "", heading).strip()
            title = alt_title if len(alt_title) > 5 else head_title
            if not title:
                title = stem.replace("-", " ")
            title = title[:1].lower() + title[1:]
            titles[stem] = f"{code} {title}"
    return titles


TITLES = diagram_titles()
all_pngs = {p.stem for p in SEQUENCE.glob("*.png")}
assert len(all_pngs) == 104
assert set(TITLES) == all_pngs, (len(TITLES), sorted(all_pngs - set(TITLES)), sorted(set(TITLES) - all_pngs))

doc = Document(SOURCE)
headings_before = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]


def find_heading(text):
    matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == text]
    assert len(matches) == 1, (text, len(matches))
    return matches[0]


def find_sd_caption(code):
    matches = [p for p in doc.paragraphs if p.style.name == "Caption" and re.search(rf"\b{re.escape(code)}\b", p.text)]
    assert len(matches) == 1, (code, len(matches))
    return matches[0]


# Move every existing Chapter 4 caption immediately below its image.
old_numbers = set()
for caption in [p for p in doc.paragraphs if p.style.name == "Caption" and re.match(r"Hình 4-\d+:", p.text)]:
    old_number = int(re.match(r"Hình 4-(\d+):", caption.text).group(1))
    image_xml = caption._p.getnext()
    assert image_xml is not None and image_xml.xpath(".//w:drawing"), (old_number, caption.text)
    image_xml.addnext(caption._p)
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.keep_together = True
    image_paragraph = next(p for p in doc.paragraphs if p._p is image_xml)
    image_paragraph.paragraph_format.keep_with_next = True
    old_numbers.add(old_number)
assert len(old_numbers) == 12


def image_before_caption(code):
    caption = find_sd_caption(code)
    previous = caption._p.getprevious()
    assert previous is not None and previous.xpath(".//w:drawing"), code
    return next(p for p in doc.paragraphs if p._p is previous)


inserted = []


def add_figure(anchor, stem):
    assert stem in TITLES, stem
    path = SEQUENCE / f"{stem}.png"
    image_paragraph = anchor.insert_paragraph_before(style="Normal")
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = Inches(0)
    image_paragraph.paragraph_format.left_indent = Inches(0)
    image_paragraph.paragraph_format.right_indent = Inches(0)
    image_paragraph.paragraph_format.space_after = Pt(4)
    image_paragraph.paragraph_format.keep_together = True
    image_paragraph.paragraph_format.keep_with_next = True
    with Image.open(path) as bitmap:
        width_px, height_px = bitmap.size
    width_inches = min(6.45, 6.35 * width_px / height_px)
    shape = image_paragraph.add_run().add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", TITLES[stem])
    caption = anchor.insert_paragraph_before(style="Caption")
    caption.add_run(f"Hình 4-0: {TITLES[stem]}")
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = False
    inserted.append(stem)


def add_many_before(anchor, stems):
    for stem in stems:
        add_figure(anchor, stem)


def find_stems(pattern):
    stems = [stem for stem in all_pngs if re.match(pattern, stem)]
    def key(stem):
        match = re.match(r"sd-(\d+)([a-z]?)", stem)
        return int(match.group(1)), match.group(2)
    return sorted(stems, key=key)


def add_group(anchor, title, intro, stems):
    heading = anchor.insert_paragraph_before(title, style="Heading 3")
    heading.paragraph_format.keep_with_next = True
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.add_run(intro)
    paragraph.paragraph_format.keep_together = True
    add_many_before(anchor, stems)


# Fill the established dynamic-view subsections around existing figures.
add_figure(image_before_caption("SD-02B"), "sd-02a-jpulse-identity")
add_many_before(image_before_caption("SD-03C"), find_stems(r"sd-03[ab]-"))
add_figure(image_before_caption("SD-06B"), "sd-06a-export-validate")
add_many_before(image_before_caption("SD-17"), find_stems(r"sd-(?:0[7-9]|1[0-6])(?:[ab])?-"))
add_figure(image_before_caption("SD-01B"), "sd-01a-jpos-identity")
add_figure(image_before_caption("SD-04A"), "sd-01c-jpos-offline")

anchor_45 = find_heading("Góc nhìn triển khai")
add_figure(anchor_45, "sd-04d-invoice-confirm")
add_group(
    anchor_45,
    "Quản lý hóa đơn điện tử",
    "Các sơ đồ SD-46–60 đi tiếp từ bước đồng bộ đơn nguồn, cấu hình, bản nháp và xem trước tới phát hành lô, theo dõi, thử lại, sổ và đối soát. Mỗi sơ đồ ghi rõ nhánh lỗi hoặc kết quả chưa xác định; trạng thái MISA và dữ liệu JPULSE cần được đối chiếu trước khi quyết định phát hành lại.",
    find_stems(r"sd-(?:4[6-9]|5\d|60)(?:[ab])?-"),
)
add_group(
    anchor_45,
    "Thông báo thời gian thực",
    "SD-05 thể hiện Web đọc và đánh dấu thông báo trên Firestore theo Security Rules. Đây là đường truy cập client được giới hạn, khác với các lệnh nghiệp vụ phải đi qua API.",
    find_stems(r"sd-05-"),
)
add_group(
    anchor_45,
    "Tổ chức, cơ sở và cấu trúc kho",
    "Các sơ đồ SD-18–35 bao phủ danh sách, tạo/sửa/xóa mềm tổ chức, cơ sở, vị trí, ô kệ, gán sản phẩm, xem và xuất dữ liệu kho cùng định mức tồn. API áp dụng quyền và phạm vi cơ sở trước khi thay đổi cấu trúc dữ liệu.",
    find_stems(r"sd-(?:1[89]|2\d|3[0-5])-"),
)
add_group(
    anchor_45,
    "Công việc, phê duyệt và kiểm kê",
    "Các sơ đồ SD-36–45 mô tả hàng công việc, quyết định phê duyệt hoặc hủy, báo cáo không phù hợp và danh sách/chi tiết phiên kiểm kê. Chúng bổ sung luồng kiểm kê SD-12–14 bằng góc nhìn của người xử lý công việc.",
    find_stems(r"sd-(?:3[6-9]|4[0-5])-"),
)
add_group(
    anchor_45,
    "Voucher marketing",
    "Các sơ đồ SD-61–80 trình bày vòng đời chiến dịch và mã voucher, tác vụ sinh mã/gia hạn, gửi email, xuất tệp và giám sát hoặc tiếp tục job. Các worker phải ghi tiến độ bền vững để retry không tạo mã hay gửi email trùng.",
    find_stems(r"sd-(?:6\d|7\d|80)-"),
)
add_group(
    anchor_45,
    "Chi phí và báo cáo chi phí",
    "Các sơ đồ SD-81–90 bao phủ nhập và sửa chi phí, Excel, chốt/mở kỳ và đối chiếu doanh thu trên báo cáo. Quyền theo cơ sở và kỳ được kiểm tra tại API trước khi ghi hoặc thay đổi trạng thái kỳ.",
    find_stems(r"sd-(?:8[1-9]|90)-"),
)

# Correct the introductory wording after switching from a selected set to the full archive.
for paragraph in doc.paragraphs:
    if paragraph.style.name == "Normal" and paragraph.text.startswith("Góc nhìn động chọn những sequence diagram"):
        paragraph.text = (
            "Góc nhìn động tập hợp toàn bộ 104 sequence diagram hiện có trong docs/architecture/sequence/, "
            "nhóm theo nghiệp vụ để tra cứu. Mỗi hình ghi lại một bước hoặc một nhánh của luồng; "
            "mã SD-A/B/C/D đánh dấu các phần nối tiếp. Mũi tên liền biểu diễn lời gọi, mũi tên đứt "
            "biểu diễn phản hồi hoặc dữ liệu đẩy."
        )
        break
else:
    raise AssertionError("Dynamic-view introductory paragraph not found")

from collections import Counter
assert len(inserted) == 95, (len(inserted), [item for item, count in Counter(inserted).items() if count > 1], sorted(all_pngs - set(inserted)))
original_stems = {
    "sd-02b-jpulse-session", "sd-02c-jpulse-mfa", "sd-03c-import-persist",
    "sd-06b-export-persist", "sd-17-transfer-receiving", "sd-01b-jpos-access",
    "sd-04a-invoice-queue", "sd-04b-invoice-worker", "sd-04c-invoice-outcome",
}
assert set(inserted).isdisjoint(original_stems)
assert set(inserted) | original_stems == all_pngs

# Number captions in the order encountered and update every existing prose cross-reference.
chapter4_start = next(i for i, p in enumerate(doc.paragraphs) if p.style.name == "Heading 1" and p.text == "Các góc nhìn kiến trúc")
chapter5_start = next(i for i, p in enumerate(doc.paragraphs) if i > chapter4_start and p.style.name == "Heading 1")
chapter4_paragraphs = doc.paragraphs[chapter4_start:chapter5_start]
captions = [p for p in chapter4_paragraphs if p.style.name == "Caption" and re.match(r"Hình 4-\d+:", p.text)]
assert len(captions) == 107, len(captions)
old_to_new = {}
for index, caption in enumerate(captions, 1):
    previous_number = int(re.match(r"Hình 4-(\d+):", caption.text).group(1))
    if previous_number in old_numbers:
        old_to_new[previous_number] = index
    title = re.sub(r"^Hình 4-\d+:\s*", "", caption.text)
    caption.text = f"Hình 4-{index}: {title}"
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = False
assert len(old_to_new) == 12, old_to_new

for paragraph in chapter4_paragraphs:
    if paragraph.style.name == "Caption":
        continue
    for run in paragraph.runs:
        updated = re.sub(
            r"Hình 4-(\d+)\b",
            lambda match: f"Hình 4-{old_to_new[int(match.group(1))]}" if int(match.group(1)) in old_to_new else match.group(0),
            run.text,
        )
        if updated != run.text:
            run.text = updated

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")][:len(headings_before)] != []
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
print("sequence_images", len(all_pngs), "chapter4_figures", len(captions), "new_headings", [p.text for p in chapter4_paragraphs if p.style.name == "Heading 3"][-6:])
