from pathlib import Path

from docx import Document


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE-CapNhat-3.1-3.3.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-3.4.1.docx"

doc = Document(SOURCE)
headings_before = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]


def heading(text: str):
    matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == text]
    assert len(matches) == 1, (text, len(matches))
    return matches[0]


def add_before(anchor, text: str, lead: str | None = None):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.paragraph_format.keep_together = True
    if lead:
        run = paragraph.add_run(lead)
        run.bold = True
    paragraph.add_run(text)


anchor = heading("API nghiệp vụ JPULSE")
add_before(anchor, "Web JPULSE là ứng dụng Next.js/React trong apps/fe-wms, cung cấp giao diện qua trình duyệt cho nhân viên và người quản trị. Các tuyến giao diện hiện có bao phủ dashboard, kho và sản phẩm, nhập xuất và điều chuyển, kiểm kê, nhân sự và chấm công, doanh thu và chi phí, hóa đơn điện tử, voucher, quản trị tài khoản/quyền cùng các chức năng cấu hình. Trang yêu cầu hóa đơn bằng token phục vụ khách hàng ở phạm vi công khai được tách khỏi khu vực quản trị. Giao diện quản trị hướng đến desktop; các luồng hiện trường cần dùng trên thiết bị di động được thiết kế theo yêu cầu Mobile Web của SRS. JPOS là hệ thống POS độc lập, không thuộc ứng dụng Web JPULSE (SRS, Mục 1.2 và 2.5).")
add_before(anchor, "Thành phần web chịu trách nhiệm điều hướng, hiển thị dữ liệu và trạng thái quy trình, thu thập đầu vào, hỗ trợ lọc/tìm kiếm, phản hồi thao tác và trình bày lỗi cho người dùng. Web có thể kiểm tra định dạng và điều kiện đơn giản để giảm thao tác sai, nhưng quy tắc nghiệp vụ, quyền thực hiện, tính hợp lệ của trạng thái và thay đổi dữ liệu có tác động quan trọng phải được API kiểm tra lại. Các màn hình dùng chung kiểu dữ liệu và trạng thái cốt lõi từ @bduck/shared-types để giữ cách diễn giải nhất quán với máy chủ.", "Trách nhiệm giao diện: ")
add_before(anchor, "Người dùng đăng nhập bằng Firebase Web SDK; web gửi ID token tới JPULSE API để thiết lập phiên và nhận thông tin tài khoản, vai trò cùng phạm vi truy cập. Khu vực dashboard kiểm tra trạng thái đăng nhập và quyền đối với trang, chuyển người chưa đăng nhập đến trang đăng nhập, hiển thị trang 403 khi không đủ quyền và áp dụng màn hình khóa MFA khi cần. Việc ẩn hoặc khóa thao tác trên web chỉ là lớp hướng dẫn người dùng; API vẫn là điểm quyết định quyền và phạm vi cơ sở cho từng yêu cầu (SRS: SEC01–SEC03).", "Xác thực và phân quyền: ")
add_before(anchor, "Lệnh tạo/sửa dữ liệu nghiệp vụ được web gửi tới JPULSE API qua HTTP/JSON để máy chủ kiểm tra đầu vào, quyền, trạng thái và audit. Một số màn hình đọc dữ liệu Firestore theo thời gian thực bằng Firebase Web SDK với truy vấn giới hạn theo cơ sở; các thao tác ghi trực tiếp chỉ áp dụng cho phạm vi hẹp được Security Rules cho phép, không thay thế API đối với mutation nghiệp vụ. Web có luồng tải tệp qua Firebase Storage SDK và luồng sử dụng URL/chính sách tải do API cấp; quyền đối với bucket và tệp phải được kiểm soát phù hợp với từng luồng (SRS: TC05–TC06).", "Giao tiếp và dữ liệu: ")
add_before(anchor, "Giao diện cần thể hiện trạng thái đang tải, đang xử lý, thành công, lỗi và chưa đồng bộ một cách rõ ràng. Với tác vụ nền, web hiển thị mã/trạng thái và tiến độ thay vì giữ kết nối chờ; lỗi nhập liệu đặt cạnh thao tác liên quan, còn lỗi hệ thống có thông điệp có thể xử lý mà không lộ chi tiết kỹ thuật. Bố cục phải phù hợp màn hình quản trị mật độ dữ liệu cao và các thao tác cảm ứng tại hiện trường; thuật ngữ, ngày giờ và trạng thái phải thống nhất với SRS, bao gồm các yêu cầu về khả năng tiếp cận và phản hồi khi mạng không ổn định (SRS: P02, USA01–USA08).", "Trải nghiệm và trạng thái: ")

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == headings_before
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
