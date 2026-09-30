from pathlib import Path

from docx import Document


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE-CapNhat-3.5.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-3.6-complete.docx"

doc = Document(SOURCE)
headings_before = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
anchors = [p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "Các góc nhìn kiến trúc"]
assert len(anchors) == 1
anchor = anchors[0]


def add(text: str, lead: str | None = None) -> None:
    paragraph = anchor.insert_paragraph_before(style="Normal")
    paragraph.paragraph_format.keep_together = True
    if lead:
        run = paragraph.add_run(lead)
        run.bold = True
    paragraph.add_run(text)


add(
    "Cách tổ chức Web JPULSE, JPULSE API và các kho dữ liệu phản ánh cả phạm vi nghiệp vụ lẫn nền tảng đang sử dụng. Web tập trung vào hiển thị và thao tác; API tập trung quy tắc, quyền và tích hợp; Firestore cùng Storage lưu trạng thái bền vững. Ranh giới này cho phép thay đổi giao diện mà không nhân bản quy tắc nghiệp vụ, đồng thời giúp nhóm phát triển xác định nơi xử lý khi bổ sung các miền kho, nhân sự, doanh thu hoặc hóa đơn. Nó cũng phù hợp ràng buộc monorepo TypeScript, Next.js/React, Node.js/Express và Firebase đã nêu ở Mục 2.7 (SRS: TC01–TC06).",
    "Tách giao diện và xử lý nghiệp vụ: ",
)
add(
    "Đặt các lệnh tạo, sửa và phê duyệt quan trọng tại API tạo một điểm kiểm soát thống nhất cho xác thực, quyền theo vai trò và phạm vi cơ sở, hợp lệ hóa dữ liệu, chuyển trạng thái và audit. Điều này cần thiết vì dữ liệu kho, số dư, hóa đơn và hồ sơ nhân sự có thể bị truy cập từ nhiều màn hình hoặc tác nhân; kiểm tra tại web không đủ để bảo vệ chúng. Đường đọc Firestore thời gian thực từ web chỉ phục vụ những nhu cầu phản hồi nhanh đã được Security Rules giới hạn. Sự kết hợp này đáp ứng yêu cầu bảo mật mà vẫn giữ trải nghiệm cập nhật kịp thời; đổi lại, quyền ở API và Rules phải được kiểm thử cùng nhau để tránh lệch chính sách (SRS: SEC01–SEC10, BC01–BC07).",
    "Ưu tiên tính đúng đắn và quyền truy cập: ",
)
add(
    "Firestore phù hợp với mô hình dữ liệu theo miền và các màn hình cần đọc cập nhật gần thời gian thực; Storage tách nội dung tệp khỏi metadata nghiệp vụ. Việc dùng dịch vụ được quản lý giảm phần hạ tầng phải tự vận hành và phù hợp định hướng triển khai trên Google Cloud. Tuy nhiên, Firestore đòi hỏi truy vấn có phạm vi, index và phân trang; thay đổi liên quan nhiều bản ghi phải dùng transaction/batch cùng cơ chế chống gửi trùng. Tệp và metadata không có một giao dịch chung nên các luồng tải tệp cần trạng thái trung gian và xử lý bù khi thất bại (SRS: TC05–TC07, SCA03, AVL03–AVL04).",
    "Lưu trữ theo trách nhiệm dữ liệu: ",
)
add(
    "API không giữ trạng thái nghiệp vụ chỉ trong bộ nhớ một instance để có thể chạy trên Cloud Run và mở rộng theo tải. Các công việc dài như phát hành hóa đơn hàng loạt hoặc xử lý voucher được theo dõi bằng trạng thái bền vững và có thể chuyển sang Cloud Tasks khi được cấu hình. API vì vậy không phải giữ yêu cầu người dùng chờ hệ thống ngoài; worker có thể thử lại sau lỗi và đối soát kết quả chưa xác định. Lựa chọn này phục vụ mục tiêu thời gian phản hồi, tính sẵn sàng và khả năng phục hồi, nhưng yêu cầu khóa chống trùng, giới hạn số lần thử, giám sát hàng đợi và thao tác xử lý thủ công cho trường hợp không thể tự khôi phục (SRS: P01–P04, SCA01–SCA05, AVL01–AVL08).",
    "Mở rộng và phục hồi tác vụ: ",
)
add(
    "Gói @bduck/shared-types giữ enum, trạng thái và payload cốt lõi nhất quán giữa web và API; các client/adapter tích hợp cô lập hợp đồng riêng của JPOS, Cityfuns/JoyWorld, MISA và những dịch vụ khác khỏi quy tắc nghiệp vụ nội bộ. Nhờ vậy, thay đổi ở một giao diện bên ngoài có phạm vi tác động rõ hơn, còn các luồng khác có thể tiếp tục khi một dịch vụ tích hợp gặp lỗi. Đánh đổi là cần quản lý phiên bản hợp đồng, cấu hình theo môi trường và kiểm thử tương thích khi triển khai web, API hoặc adapter vào những thời điểm khác nhau. Các quyết định và rủi ro cụ thể được đánh giá thêm ở Chương 7 (SRS: SCA06, TC08–TC13, COM06).",
    "Giảm phụ thuộc khi thay đổi: ",
)

assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == headings_before
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
