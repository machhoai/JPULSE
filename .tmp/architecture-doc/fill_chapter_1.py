from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\Github\bduck-system")
SOURCE = ROOT / "SAD-JPULSE.docx"
OUTPUT = ROOT / ".tmp" / "architecture-doc" / "SAD-JPULSE-chapter-1.docx"

doc = Document(SOURCE)


def heading(text: str):
    matches = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip() == text]
    assert len(matches) == 1, (text, len(matches))
    return matches[0]


def add_before(anchor, text: str, lead: str | None = None):
    paragraph = anchor.insert_paragraph_before(style="Normal")
    if lead:
        run = paragraph.add_run(lead)
        run.bold = True
    paragraph.add_run(text)
    return paragraph


anchor = heading("Phạm vi hệ thống")
add_before(anchor, "Tài liệu kiến trúc phần mềm này mô tả cách tổ chức hệ thống JPULSE để đáp ứng phạm vi và các yêu cầu đã xác định trong tài liệu SRS-JPULSE. Tài liệu xác định ranh giới hệ thống, trách nhiệm của các thành phần, cách trao đổi và lưu trữ dữ liệu, các luồng xử lý quan trọng và mô hình triển khai.")
add_before(anchor, "Tài liệu là cơ sở cho nhóm thiết kế và phát triển triển khai nhất quán giữa giao diện web, API, dữ liệu và các tích hợp; đồng thời hỗ trợ nhóm kiểm thử, vận hành và bảo trì đánh giá tác động khi hệ thống thay đổi. Yêu cầu nghiệp vụ chi tiết, quy tắc nghiệp vụ và tiêu chí chấp nhận tiếp tục được quản lý trong SRS; các chương sau của tài liệu này tập trung giải thích các quyết định kiến trúc tương ứng.")

anchor = heading("Đối tượng đọc tài liệu")
add_before(anchor, "JPULSE là nền tảng quản trị vận hành của Joy World, bao gồm quản lý kho và tồn kho, nhân sự và chấm công, doanh thu và chi phí nội bộ, hóa đơn điện tử, quản trị truy cập, thiết bị POS và chiến dịch voucher. Phạm vi kiến trúc của tài liệu bao trùm các thành phần phần mềm phục vụ những nhóm chức năng đó cùng các quan hệ giữa chúng.")
add_before(anchor, "Các thành phần được xem xét gồm giao diện web dùng trên máy tính và thiết bị di động, API xử lý nghiệp vụ, hợp đồng dữ liệu dùng chung, dịch vụ xác thực, kho dữ liệu và lưu trữ tệp. Tài liệu cũng xét các điểm kết nối với JPOS, hệ thống ERP cũ, MISA meInvoice và dịch vụ email Brevo; Firebase và Google Cloud là các dịch vụ nền tảng được SRS xác định cho dữ liệu, tệp và môi trường vận hành.")
add_before(anchor, "Ranh giới của tài liệu là kiến trúc hệ thống JPULSE và các giao diện tích hợp tại phía JPULSE. Đặc tả từng màn hình, use case, trường dữ liệu, endpoint và quy tắc nghiệp vụ chi tiết thuộc SRS hoặc tài liệu chuyên biệt. Theo phạm vi SRS, tính lương, kế toán tổng hợp và cổng thương mại điện tử cho khách hàng không nằm trong phạm vi chức năng của JPULSE.")

anchor = heading("Thuật ngữ và chữ viết tắt")
add_before(anchor, "Tài liệu dành cho nhóm thiết kế và phát triển, nhóm kiểm thử, người quản trị và vận hành hệ thống, cùng các bên phụ trách nghiệp vụ cần hiểu ranh giới và tác động của thay đổi. Mỗi nhóm sử dụng tài liệu theo mối quan tâm sau:")
add_before(anchor, " dùng để xác định trách nhiệm của giao diện, API, dữ liệu và tích hợp; làm căn cứ thiết kế chi tiết và sửa đổi hệ thống.", "Nhóm thiết kế và phát triển:")
add_before(anchor, " dùng để xây dựng các kịch bản kiểm thử luồng liên thành phần, phân quyền, tích hợp và xử lý lỗi dựa trên kiến trúc đã mô tả.", "Nhóm kiểm thử:")
add_before(anchor, " dùng để cấu hình môi trường, quản lý quyền truy cập, theo dõi sự cố, sao lưu và khôi phục dữ liệu.", "Người quản trị và vận hành:")
add_before(anchor, " dùng để kiểm tra sự phù hợp giữa giải pháp kỹ thuật với phạm vi, yêu cầu chất lượng và các phụ thuộc nghiệp vụ trong SRS.", "Ban quản lý và các bộ phận nghiệp vụ:")

caption = next(p for p in doc.paragraphs if p.style.name == "Caption" and p.text.strip() == "Bảng 1-1: Thuật ngữ và chữ viết tắt")
add_before(caption, "Các thuật ngữ dưới đây được dùng thống nhất trong tài liệu; tên dịch vụ và nhóm chức năng tuân theo cách gọi trong SRS-JPULSE.")

terms = [
    ("J-PULSE / JPULSE", "Platform for Unified Leadership, Service & Enterprise of Joy World", "Hệ thống quản trị các nghiệp vụ vận hành được mô tả trong SRS.", ""),
    ("SRS", "Software Requirements Specification", "Tài liệu đặc tả yêu cầu phần mềm của JPULSE.", ""),
    ("SAD", "Software Architecture Document", "Tài liệu mô tả kiến trúc phần mềm của JPULSE.", ""),
    ("WMS", "Warehouse Management System", "Nhóm chức năng quản lý kho, tồn kho và luân chuyển hàng hóa.", ""),
    ("JPOS", "Hệ thống bán hàng JPOS", "Hệ thống POS/kênh bán hàng trao đổi dữ liệu với JPULSE.", ""),
    ("API", "Application Programming Interface", "Giao diện trao đổi dữ liệu giữa các thành phần hoặc hệ thống.", ""),
    ("RBAC", "Role-Based Access Control", "Cơ chế phân quyền theo vai trò người dùng.", ""),
    ("Office Scope", "Phạm vi cơ sở", "Giới hạn dữ liệu và thao tác theo cơ sở được phân công.", ""),
    ("MFA", "Multi-Factor Authentication", "Xác thực người dùng bằng nhiều yếu tố.", ""),
    ("BOM", "Bill of Materials", "Định mức vật tư gắn với sản phẩm.", ""),
    ("ATP", "Available to Promise", "Số lượng có thể đáp ứng theo chính sách tồn kho.", ""),
    ("GCP", "Google Cloud Platform", "Nền tảng đám mây được SRS nêu cho môi trường vận hành.", ""),
    ("Firestore", "Cloud Firestore", "Kho dữ liệu dùng trong hệ sinh thái Firebase.", ""),
    ("meInvoice", "MISA meInvoice", "Dịch vụ phát hành và quản lý hóa đơn điện tử tích hợp với JPULSE.", ""),
]

table = doc.tables[0]
assert [cell.text for cell in table.rows[0].cells] == ["Thuật ngữ", "Tên đầy đủ", "Ý nghĩa", "Ghi chú"]
white_row = deepcopy(table.rows[1]._tr)
tint_row = deepcopy(table.rows[2]._tr)
for row in list(table.rows)[3:]:
    table._tbl.remove(row._tr)
for i in range(2, len(terms)):
    table._tbl.append(deepcopy(white_row if i % 2 == 0 else tint_row))
for row, values in zip(table.rows[1:], terms):
    for cell, value in zip(row.cells, values):
        paragraph = cell.paragraphs[0]
        for run in paragraph.runs:
            run.text = ""
        run = paragraph.add_run(value)
        run.font.size = Pt(9.5)

anchor = heading("Bối cảnh và yêu cầu kiến trúc")
add_before(anchor, "Nguồn tham chiếu chính là tài liệu SRS-JPULSE.docx, phiên bản 1.0. Chương 1 của SRS xác định phạm vi và thuật ngữ; Chương 2 mô tả bối cảnh, các bên liên quan, môi trường vận hành và ràng buộc; Chương 5 trình bày dữ liệu; Chương 6 trình bày tích hợp; Chương 7 xác định yêu cầu phi chức năng.")
add_before(anchor, "Khi thông tin trong tài liệu kiến trúc cần đối chiếu về chức năng hoặc phạm vi nghiệp vụ, áp dụng SRS-JPULSE làm tài liệu yêu cầu gốc. Đặc tả chi tiết của từng API, mô hình dữ liệu và cấu hình triển khai được tham chiếu tại các chương tương ứng của tài liệu kiến trúc khi những phần đó được hoàn thiện.")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
