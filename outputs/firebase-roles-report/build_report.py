import json, math
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone, timedelta
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).parent
projects=json.loads((ROOT/'roles-snapshot.json').read_text(encoding='utf-8'))
registry=json.loads((ROOT/'registry.json').read_text(encoding='utf-8'))
groups=sorted(registry['groups'], key=lambda g:g['order'])
perms=registry['permissions']; lookup={p['key']:p for p in perms}
groupmap={g['id']:g['label']['vi'] for g in groups}
prod=next(p for p in projects if p['project']=='jw-system-f2104')
test=next(p for p in projects if p['project']=='test-jw-system')
active=[r for r in prod['roles'] if r.get('is_deleted') is False]
byid={r['id']:r for r in active}
def font(n,bold=False):return ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf',n)
def box(draw, xy, text, sub=None):
    draw.rounded_rectangle(xy,radius=15,fill='#edf3fa',outline='#6686a6',width=3)
    x0,y0,x1,y1=xy
    draw.text(((x0+x1)/2,(y0+y1)/2-(16 if sub else 0)),text,font=font(30,True),fill='#172b40',anchor='mm')
    if sub:draw.text(((x0+x1)/2,(y0+y1)/2+28),sub,font=font(23),fill='#384e62',anchor='mm')
def diagram():
    im=Image.new('RGB',(1800,1080),'white');d=ImageDraw.Draw(im)
    names={r['name']:r for r in active}
    coords={'SUPER_ADMIN':(30,30,520,135),'Đối tác Trung Quốc':(630,30,1120,135),'Giám đốc':(1230,30,1770,135),
    'Marketing':(30,370,400,475),'Quản lý nhân sự':(450,370,850,475),'Kế toán trưởng':(900,370,1300,475),'Cửa hàng trưởng':(1350,370,1770,475),
    'Kế toán':(900,690,1300,795),'Quản lý cửa hàng':(1270,900,1700,1005),'Thu ngân':(1350,690,1770,795)}
    for r in active:
        if r.get('parent_id') in byid:
            parent=byid[r['parent_id']]['name'];a=coords[parent];b=coords[r['name']]
            start=((a[0]+a[2])/2,a[3]);end=((b[0]+b[2])/2,b[1]);mid=(start[1]+end[1])/2
            if r['name']=='Quản lý cửa hàng':
                d.line([start,(1788,start[1]+30),(1788,850),(1485,850),end],fill='#6686a6',width=4)
            else:d.line([start,(start[0],mid),(end[0],mid),end],fill='#6686a6',width=4)
    for name,xy in coords.items():
        r=names[name];n=sum(v is True for v in r.get('permissions',{}).values())
        box(d,xy,name,f'{n} khóa bật'+('  |  Có *' if r['permissions'].get('*') is True else ''))
    im.save(ROOT/'role-tree.png')
def group_diagram():
    im=Image.new('RGB',(1800,1510),'white');d=ImageDraw.Draw(im)
    box(d,(570,15,1230,120),'Danh mục quyền',f'{len(perms)} khóa trong {len(groups)} nhóm')
    for i,g in enumerate(groups):
        col=i%3;row=i//3;x=25+col*600;y=190+row*145
        d.line([(900,120),(900,150),(x+275,150),(x+275,y)],fill='#a4b7ca',width=2)
    for i,g in enumerate(groups):
        col=i%3;row=i//3;x=25+col*600;y=190+row*145
        n=sum(p['group']==g['id'] for p in perms)
        box(d,(x,y,x+550,y+110),g['label']['vi'],f"{g['id']}  |  {n} khóa")
    im.save(ROOT/'permission-groups.png')
diagram();group_diagram()
doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.65);sec.right_margin=Cm(1.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(6)
    for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(13)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Trang ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(t='',style=None):return doc.add_paragraph(t,style)
def h(t,l=1):return doc.add_heading(t,l)
def page():doc.add_page_break()
def table(headers,rows,widths,size=9):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,s in zip(t.rows[0].cells,headers):c.text=s
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if ri==0 or ri%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'203B56' if ri==0 else 'F2F5F8');tcpr.append(sh)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1
                if ri==0:pp.paragraph_format.keep_with_next=True
                for run in pp.runs:
                    run.font.size=Pt(size);run.bold=ri==0
                    if ri==0:run.font.color.rgb=RGBColor(255,255,255)
    p().paragraph_format.space_after=Pt(1)
    return t
def enabled(r):return [k for k,v in r.get('permissions',{}).items() if v is True]
def disabled(r):return [k for k,v in r.get('permissions',{}).items() if v is False]
def role_order(roles):
    out=[];seen=set()
    def walk(r):
        if r['id'] in seen:return
        seen.add(r['id']);out.append(r)
        for c in sorted([x for x in roles if x.get('parent_id')==r['id']],key=lambda x:x['name']):walk(c)
    for r in roles:
        if not r.get('parent_id'):walk(r)
    for r in roles:walk(r)
    return out

doc.add_heading('Sơ đồ cấp bậc và quyền của các role',0)
p('Hệ thống Bduck trên Firebase', 'Subtitle')
when=datetime.fromisoformat(prod['read_time']).astimezone(timezone(timedelta(hours=7))).strftime('%d/%m/%Y lúc %H:%M:%S')
p(f'Dữ liệu được đọc ngày {when} theo giờ Việt Nam.')
p('Môi trường chính jw-system-f2104 có 10 role đang hoạt động và 1 role đã xóa. Danh mục của ứng dụng có 146 khóa quyền trong 26 nhóm. Tài liệu trình bày cây role theo parent_id, cây nhóm quyền và danh sách quyền đang bật cho từng role; môi trường thử nghiệm được trình bày riêng ở cuối.')
h('Cách đọc phân quyền',2)
for text in [
'Quan hệ cha–con giữa các role biểu thị cấu trúc quản lý. Mã kiểm tra quyền sử dụng permissions của chính role, không tự động cộng quyền từ role cha.',
'Giá trị true là quyền được bật trực tiếp. Giá trị false và khóa không có trong permissions đều không cấp quyền từ role đó. Nếu một người có nhiều role hợp lệ, các quyền true được hợp nhất theo phạm vi cơ sở.',
'Khóa * cấp toàn bộ quyền trong phạm vi được gán. Một giá trị false riêng lẻ không phủ định quyền do * hoặc role khác cấp. Tên ADMIN hay SUPER_ADMIN tự nó không cấp quyền.',
'Quyền toàn hệ thống cần role có * = true và lượt gán hợp lệ với warehouse_id = null. Quyền tại cơ sở còn phụ thuộc loại cơ sở, trạng thái tài khoản, thời hạn và lượt gán.',
'Các số lượng trong báo cáo là số khóa được lưu, không phải số người dùng hoặc số cơ sở được phép truy cập. Danh sách role không xác nhận quyền hiệu lực của từng tài khoản.']:
    p(text)
h('Các điểm cần lưu ý',2)
p('Giám đốc và SUPER_ADMIN đang có * = true ở môi trường chính. Giám đốc đồng thời lưu stock_counts.view = false; giá trị này không chặn quyền xem kiểm đếm khi * đang có hiệu lực. Role con Kế toán có 70 khóa bật, nhiều hơn 53 khóa của Kế toán trưởng; đây là cấu hình trực tiếp, không phải quyền kế thừa.')
p('Môi trường thử nghiệm có hai role có * = true là ADMIN và Quản lý kho. Tất cả khóa quyền đang lưu trong các role đều có trong danh mục ứng dụng.')
page();h('Sơ đồ cấp bậc role môi trường chính')
p('Mũi nối biểu thị role cha → role con theo trường parent_id. Ba role gốc độc lập là SUPER_ADMIN, Đối tác Trung Quốc và Giám đốc. Không có mối nối nào biểu thị kế thừa quyền.')
doc.add_picture(str(ROOT/'role-tree.png'),width=Cm(17.7))
h('Tổng hợp các role đang hoạt động',2)
table(['Role','Role cha','Bật','Tắt','Có *'],[(r['name'],byid.get(r.get('parent_id'),{}).get('name','Không có'),len(enabled(r)),len(disabled(r)),'Có' if r['permissions'].get('*') is True else 'Không') for r in role_order(active)],[5,5,2,2,3.7])
p('Role Giám đốc đã xóa có ID khác với role Giám đốc đang hoạt động và được tách riêng trong phụ lục.')
page();h('Sơ đồ nhóm quyền của hệ thống')
p('Cấu trúc danh mục là Hệ thống → Nhóm quyền → Khóa quyền. Đây là phân nhóm chức năng để tra cứu; các quyền cùng nhóm không tự cấp lẫn nhau. Các lá của từng nhóm được liệt kê đầy đủ trong phần kế tiếp.')
doc.add_picture(str(ROOT/'permission-groups.png'),width=Cm(17.7))
p('Danh mục gồm 145 khóa chức năng và 1 khóa toàn quyền *. Số lượng này lấy từ permissionRegistry.ts trong mã nguồn tại thời điểm lập báo cáo.')
page();h('Danh mục đầy đủ các khóa quyền')
p('Mỗi tiêu đề bên dưới là một nhánh trong sơ đồ nhóm quyền. Bảng liệt kê toàn bộ các khóa lá và tên hiển thị tiếng Việt của nhánh đó. Giữ nguyên tên trong danh mục ứng dụng, kể cả tên chưa có dấu.')
for g in groups:
    items=[x for x in perms if x['group']==g['id']]
    h(f"{g['label']['vi']} gồm {len(items)} khóa",2)
    table(['Khóa quyền','Tên hiển thị'],[(x['key'],x['label']['vi']) for x in items],[8.5,9.2],9)

def role_section(r,roles,prefix=''):
    page();h(prefix+r['name'])
    parent=next((x['name'] for x in roles if x['id']==r.get('parent_id')), 'Không có')
    p('ID role: '+r['id']);p(f"Role cha: {parent}. Trạng thái: {'Đã xóa' if r.get('is_deleted') else 'Đang hoạt động'}.")
    if r.get('description'):p('Mô tả: '+r['description'])
    if r.get('is_deleted'):p('Role này đã bị xóa mềm và không được dùng để cấp quyền hiện hành. Các khóa dưới đây là cấu hình được lưu lại để tra cứu lịch sử.')
    p(f"Quyền lưu trực tiếp: {len(enabled(r))} khóa bật và {len(disabled(r))} khóa tắt.")
    if r.get('permissions',{}).get('*') is True:p('Role có * = true. Các khóa bật còn lại là cấu hình lưu bổ sung; * bao phủ toàn bộ danh mục trong phạm vi gán, bao gồm cả các khóa không được liệt kê riêng. Các khóa false dưới đây không tạo ngoại lệ từ chối đối với *.')
    h('Các quyền được cấp trực tiếp',2)
    rows=[]
    for g in groups:
        ks=[x for x in perms if x['group']==g['id'] and r.get('permissions',{}).get(x['key']) is True]
        if ks:rows.append((g['label']['vi']+'\n'+str(len(ks))+' khóa','\n'.join(x['key'] for x in ks)))
    table(['Nhóm quyền','Khóa đang bật'],rows,[4.5,13.2],9.5)
    if disabled(r):
        h('Các khóa được lưu với giá trị false',2)
        p('; '.join(sorted(disabled(r)))+'.')
    else:p('Không có khóa nào được lưu với giá trị false.')
    p('Các khóa còn lại trong danh mục không được lưu riêng trên role này.')

for r in role_order(active):role_section(r,active)
for r in prod['roles']:
    if r.get('is_deleted'):role_section(r,prod['roles'],'Phụ lục role đã xóa ')
page();h('Phụ lục môi trường thử nghiệm')
p('Firebase project test-jw-system có 4 role đang hoạt động. Cả 4 role đều là role gốc do parent_id rỗng hoặc không tồn tại. Dữ liệu ở đây được tách khỏi môi trường chính.')
table(['Role gốc','Khóa bật','Khóa tắt','Có *'],[(r['name'],len(enabled(r)),len(disabled(r)),'Có' if r['permissions'].get('*') is True else 'Không') for r in test['roles']],[7.7,3,3,4])
p('Sơ đồ cấp bậc: test-jw-system → ADMIN; Kế toán; Test; Quản lý kho. Không có quan hệ cha–con giữa các role này.')
for r in role_order(test['roles']):role_section(r,test['roles'],'Thử nghiệm ')
page();h('Nguồn dữ liệu và quy tắc đối chiếu')
p('Dữ liệu role được đọc trực tiếp từ collection roles trên hai Firebase project jw-system-f2104 và test-jw-system ngày 26/10/2026. Báo cáo giữ nguyên ID tài liệu, name, parent_id, permissions và is_deleted khi đối chiếu; không chứa khóa bí mật hoặc thông tin tài khoản cá nhân.')
table(['Nguồn','Nội dung đối chiếu'],[
('Firebase roles','Cây role, trạng thái, khóa bật và khóa tắt'),
('packages/shared-types/src/permissionRegistry.ts','146 khóa quyền, tên tiếng Việt, 26 nhóm'),
('apps/be-wms/src/services/rolePermissionUtils.ts','Quyền role lấy từ permissions được lưu; tên role không cấp quyền'),
('apps/be-wms/src/services/authorization/accessContextBuilder.ts','Hợp nhất quyền true và điều kiện wildcard toàn hệ thống'),
('apps/be-wms/src/services/authorization/authorizationService.ts','Kiểm tra khóa * hoặc khóa cụ thể cùng giới hạn loại cơ sở'),
('apps/be-wms/src/services/roleService.ts','Tạo và kiểm tra quan hệ parent_id'),
],[8.5,9.2],9)
p('Kết quả đối chiếu: 11 tài liệu role ở môi trường chính, trong đó 10 chưa xóa; 4 tài liệu ở môi trường thử nghiệm, đều chưa xóa. Không phát hiện khóa ngoài danh mục, role con đang hoạt động tham chiếu cha bị thiếu hoặc chu trình trong cây role chính.')
out=ROOT/'So-do-phan-quyen-Firebase-Bduck.docx';doc.save(out)
print(out)
