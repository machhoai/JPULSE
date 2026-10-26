import json
from pathlib import Path
from collections import defaultdict
from datetime import datetime,timezone,timedelta,date
from copy import deepcopy
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).parent
src=ROOT/'Bao-cao-phan-quyen-danh-cho-lanh-dao.docx'
data=json.loads((ROOT/'active-role-members-snapshot.json').read_text(encoding='utf-8'))
now=datetime.fromisoformat(data['read_time'].replace('Z','+00:00'))
local=now.astimezone(timezone(timedelta(hours=7)))
def boundary(value,start):
    if value is None:return True
    if isinstance(value,str):
        try:
            if len(value)==10:
                d=date.fromisoformat(value)
                return local.date()>=d if start else local.date()<=d
            d=datetime.fromisoformat(value.replace('Z','+00:00'))
            return now>=d if start else now<=d
        except ValueError:return False
    if isinstance(value,dict):
        seconds=value.get('_seconds',value.get('seconds'))
        if seconds is not None:
            t=datetime.fromtimestamp(seconds,tz=timezone.utc)
            return now>=t if start else now<=t
    return False
def valid_window(r,optional=False):
    if not optional and ('valid_from' not in r or 'valid_until' not in r):return False
    return boundary(r.get('valid_from'),True) and boundary(r.get('valid_until'),False)
roles={r['id']:r for r in data['roles'] if r.get('is_deleted') is False and r.get('is_active',True) is True and r.get('status','ACTIVE')=='ACTIVE' and valid_window(r,True)}
users={u['id']:u for u in data['users'] if u.get('status')=='ACTIVE' and u.get('is_deleted') is False}
facilities={w['id']:w for w in data['facilities']}
members=defaultdict(lambda:defaultdict(set));included=[];excluded=[]
for a in data['assignments']:
    if a.get('user_id') not in users or a.get('role_id') not in roles or a.get('is_active') is not True or a.get('is_deleted',False) is not False or a.get('scope_origin','DIRECT') not in ['DIRECT','LEGACY_DIRECT'] or not valid_window(a):
        excluded.append(a['id']);continue
    if 'warehouse_id' not in a or (a['warehouse_id'] is not None and a['warehouse_id'] not in facilities):
        excluded.append(a['id']);continue
    scope='Toàn hệ thống' if a['warehouse_id'] is None else facilities[a['warehouse_id']]['name']
    members[a['role_id']][a['user_id']].add(scope);included.append(a['id'])
def rn(r):return 'Quản trị hệ thống' if r['name']=='SUPER_ADMIN' else r['name']
byname={rn(r):r for r in roles.values()}
doc=Document(src)
original_paragraphs=[p.text for p in doc.paragraphs]
original_tables=[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]

def para_after(anchor,text,style='Normal'):
    p=doc.add_paragraph(text,style);anchor.addnext(p._p);return p._p
def add_table(anchor,rows):
    t=doc.add_table(rows=1,cols=3);t.autofit=False
    widths=[5.2,4.4,8.1]
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,s in zip(t.rows[0].cells,['Họ và tên','Tài khoản','Cơ sở được gán']):c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=s
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c,w in zip(row.cells,widths):
            c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bd.append(el)
            pr.append(bd)
            marg=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');marg.append(el)
            pr.append(marg)
            if i==0 or i%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'203B56' if i==0 else 'F2F5F8');pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(1);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
                if i==0:p.paragraph_format.keep_with_next=True
                for run in p.runs:
                    run.font.name='Arial';run.font.size=Pt(9.5);run.bold=i==0
                    run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    anchor.addnext(t._tbl);return t._tbl
def insert_members(anchor,name):
    r=byname[name];m=members[r['id']]
    anchor=para_after(anchor,f'Người dùng đang hoạt động thuộc vai trò này gồm {len(m)} tài khoản','Heading 2')
    if not m:return para_after(anchor,'Chưa có tài khoản đang hoạt động được gán vai trò còn hiệu lực.')
    rows=[]
    for uid in sorted(m,key=lambda x:users[x]['full_name'].casefold()):
        u=users[uid]
        rows.append((u.get('full_name') or 'Chưa có họ tên',u.get('username') or u.get('email') or 'Chưa có tên tài khoản','\n'.join(sorted(m[uid]))))
    return add_table(anchor,rows)

# Bổ sung mô tả ngắn ngay sau bảng quy mô hiện có.
info=doc.tables[0]._tbl
unique={uid for rm in members.values() for uid in rm}
info=para_after(info,f"Danh sách người dùng được cập nhật ngày {local.strftime('%d/%m/%Y lúc %H:%M')} theo giờ Việt Nam. Có {len(unique)} tài khoản đang hoạt động có ít nhất một lượt gán vai trò còn hiệu lực trong hệ thống chính.")
para_after(info,'Chỉ liệt kê tài khoản đang hoạt động, chưa xóa và có lượt gán vai trò đang bật, còn thời hạn. Một người có thể xuất hiện ở nhiều vai trò; nhiều cơ sở trong cùng vai trò được gộp trên một dòng. Cột cơ sở thể hiện nơi được gán trực tiếp trong hệ thống.')

# Các vai trò đã có phần chi tiết: bổ sung danh sách ngay trước bảng quyền hiện có.
for paragraph in list(doc.paragraphs):
    if paragraph.text.startswith('Quyền của '):
        name=paragraph.text[len('Quyền của '):]
        if name not in byname:continue
        node=paragraph._p.getnext();last=paragraph._p
        while node is not None and node.tag!=qn('w:tbl'):
            if node.tag==qn('w:p'):
                # Dừng trước tiêu đề tiếp theo nếu cấu trúc tài liệu thay đổi.
                ps=node.find('w:pPr/w:pStyle',node.nsmap)
                if ps is not None and ps.get(qn('w:val')) in ['Heading1','Heading2']:break
            last=node;node=node.getnext()
        end=insert_members(last,name)
        para_after(end,'Các chức năng được cấp','Heading 2')

# Hai vai trò toàn quyền chỉ có trong bảng tổng quan nên thêm danh sách ngay sau bảng này.
overview=next(t for t in doc.tables if t.cell(0,0).text=='Vai trò' and t.cell(0,1).text=='Được làm gì')
anchor=overview._tbl
for name in ['Quản trị hệ thống','Giám đốc']:
    anchor=para_after(anchor,'Người dùng thuộc vai trò '+name,'Heading 2')
    anchor=insert_members(anchor,name)

out=ROOT/'Bao-cao-phan-quyen-bo-sung-nguoi-dung.docx'
doc.save(out)
check=Document(out)
# Mọi đoạn và bảng cũ phải còn nguyên nội dung và thứ tự.
idx=0
for p in check.paragraphs:
    if idx<len(original_paragraphs) and p.text==original_paragraphs[idx]:idx+=1
assert idx==len(original_paragraphs),'An original paragraph was changed or removed'
idx=0
for t in check.tables:
    contents=[[c.text for c in row.cells] for row in t.rows]
    if idx<len(original_tables) and contents==original_tables[idx]:idx+=1
assert idx==len(original_tables),'An original table was changed or removed'
summary={rn(r):len(members[r['id']]) for r in roles.values()}
(ROOT/'active-role-members-audit.json').write_text(json.dumps({'unique_active_users_with_valid_roles':len(unique),'active_users_without_valid_roles':[users[u]['full_name'] for u in users if u not in unique],'included_assignments':included,'excluded_assignments':excluded,'counts':summary},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(out),'unique_users':len(unique),'counts':summary,'preserved_paragraphs':len(original_paragraphs),'preserved_tables':len(original_tables)},ensure_ascii=False))
